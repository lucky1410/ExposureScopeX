from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import re
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Literal
from uuid import UUID

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .serialization import json_safe


EVALUATOR_REPORT_VERSION = "1.6"
ReportFormat = Literal["docx", "pdf"]
REPORT_MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
REPORT_EXTENSIONS = {"docx": "docx", "pdf": "pdf"}


class EvaluationReportError(RuntimeError):
    pass


def _canonical_json(value: object) -> bytes:
    return json.dumps(
        json_safe(value), sort_keys=True, separators=(",", ":"), ensure_ascii=True
    ).encode("utf-8")


def _sha256(value: object) -> str:
    return hashlib.sha256(_canonical_json(value)).hexdigest()


def _safe_text(value: object, limit: int = 2000) -> str:
    return str(value if value is not None else "-").replace("\x00", "")[:limit]


def _percent(value: object) -> str:
    if isinstance(value, (int, float)):
        return f"{value:.1%}"
    return "Not measured"


def _number(value: object) -> str:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return f"{value:.3f}"
    return _safe_text(value)


def _usd(value: object) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"${value:,.6f}"
    return "Not measured"


def _status(value: object) -> str:
    if value is True:
        return "PASS"
    if value is False:
        return "FAIL"
    return "NOT MEASURABLE"


def _format_timestamp(value: object) -> str:
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    except (TypeError, ValueError):
        return _safe_text(value)


def _label(value: str) -> str:
    return value.replace("_", " ").replace("@", " at ").title()


def evaluation_snapshot(evaluation: dict) -> dict:
    """Return the immutable, reportable portion of an evaluation record."""
    fields = (
        "id",
        "name",
        "evaluated_agent_id",
        "evaluated_agent_version",
        "evaluator_agent_id",
        "evaluator_version",
        "dataset_version",
        "project_key",
        "dataset",
        "input_manifest",
        "metrics",
        "release_decision",
        "created_at",
    )
    return {field: json_safe(evaluation.get(field)) for field in fields}


def evaluation_source_sha256(evaluation: dict) -> str:
    return _sha256(evaluation_snapshot(evaluation))


def _input_manifest_sha256(evaluation: dict) -> str:
    return _sha256(evaluation.get("input_manifest") or {})


def _metrics_sha256(evaluation: dict) -> str:
    return _sha256(evaluation.get("metrics") or {})


def _metric_rows(metrics: dict) -> list[list[str]]:
    return [
        ["Release decision", _safe_text(metrics.get("release_decision", "inconclusive")).upper()],
        ["Overall score", _number(metrics.get("overall_score"))],
        ["Measurement coverage", _percent(metrics.get("measurement_coverage"))],
        ["Required dimensions", ", ".join(_label(item) for item in metrics.get("required_dimensions", [])) or "None"],
        ["Missing required dimensions", ", ".join(_label(item) for item in metrics.get("missing_required_dimensions", [])) or "None"],
    ]


def _classification_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("classification") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    rows = [
        ["Sample size", _number(item.get("sample_size"))],
        ["Ground-truth classes", _number(item.get("ground_truth_class_count"))],
        ["Accuracy", _percent(item.get("accuracy"))],
        ["Macro precision", _percent(item.get("macro_precision"))],
        ["Macro recall", _percent(item.get("macro_recall"))],
        ["Macro F1", _percent(item.get("macro_f1"))],
        ["False-positive rate", _percent(item.get("macro_false_positive_rate"))],
        ["False-negative rate", _percent(item.get("macro_false_negative_rate"))],
    ]
    return rows


def _confusion_rows(metrics: dict) -> list[list[str]]:
    matrix = (metrics.get("classification") or {}).get("confusion_matrix") or {}
    labels = list((metrics.get("classification") or {}).get("labels") or [])
    if not matrix or not labels:
        return [["Not measurable", "No labelled confusion matrix was supplied."]]
    rows: list[list[str]] = []
    for expected in labels:
        for predicted in labels:
            rows.append([_safe_text(expected), _safe_text(predicted), _number((matrix.get(expected) or {}).get(predicted, 0))])
    return rows


def _confidence_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("confidence") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    return [
        ["Correctness Brier score", _number(item.get("correctness_brier_score"))],
        ["Expected calibration error", _percent(item.get("expected_calibration_error"))],
    ]


def _grounding_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("groundedness") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    return [
        ["Claims reviewed", _number(item.get("claim_count"))],
        ["Supported claim rate", _percent(item.get("supported_claim_rate"))],
        ["Unsupported claim rate", _percent(item.get("unsupported_claim_rate"))],
        ["Citation validity", _percent(item.get("citation_validity_rate"))],
        ["Evidence integrity", _percent(item.get("evidence_integrity_rate"))],
        ["Unsupported claim IDs", ", ".join(item.get("unsupported_claim_ids") or []) or "None"],
    ]


def _security_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("security") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    confusion = item.get("detection_confusion") or {}
    return [
        ["Security cases", _number(item.get("case_count"))],
        ["Attack outcome accuracy", _percent(item.get("attack_outcome_accuracy"))],
        ["Attack success rate", _percent(item.get("attack_success_rate"))],
        ["Detection rate", _percent(item.get("detection_rate"))],
        ["False detection rate", _percent(item.get("false_detection_rate"))],
        ["Evidence coverage", _percent(item.get("evidence_coverage"))],
        ["Detection TP / FP / FN / TN", " / ".join(_number(confusion.get(key, 0)) for key in ("tp", "fp", "fn", "tn"))],
    ]


def _trajectory_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("trajectory") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    return [
        ["Trajectory score", _percent(item.get("score"))],
        ["Milestone coverage", _percent(item.get("milestone_coverage"))],
        ["Action efficiency", _percent(item.get("action_efficiency"))],
        ["Policy compliant", "Yes" if item.get("policy_compliant") else "No"],
        ["Missing milestones", ", ".join(item.get("missing_milestones") or []) or "None"],
        ["Scope violations", ", ".join(item.get("scope_violations") or []) or "None"],
        ["Tool misuse", ", ".join(item.get("tool_misuse_events") or []) or "None"],
    ]


def _optional_dimension_rows(metrics: dict) -> list[list[str]]:
    rows: list[list[str]] = []
    for key in ("rag", "robustness", "judge_agreement", "reproducibility", "cost_efficiency"):
        item = metrics.get(key) or {}
        status = item.get("measurement_status", "not_measurable")
        if status == "measured":
            measurable = [
                f"{_label(name)}: {_percent(value) if isinstance(value, float) else _number(value)}"
                for name, value in item.items()
                if name not in {"measurement_status", "definition", "bins", "reason"}
                and not isinstance(value, (dict, list))
            ]
            rows.append([_label(key), "Measured", "; ".join(measurable) or "Measured"])
        else:
            rows.append([_label(key), "Not measurable", _safe_text(item.get("reason"))])
    return rows


def _rag_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("rag") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    return [
        ["Context precision", _percent(item.get("context_precision"))],
        ["Recall at K", _percent(item.get("recall_at_k"))],
        ["Mean reciprocal rank", _number(item.get("mean_reciprocal_rank"))],
        ["Faithfulness", _percent(item.get("faithfulness"))],
        ["Citation validity", _percent(item.get("citation_validity"))],
        ["Uncited relevant documents", ", ".join(item.get("uncited_relevant_documents") or []) or "None"],
    ]


def _robustness_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("robustness") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    counts = item.get("by_variation_type") or {}
    variation_summary = "; ".join(
        f"{_label(str(kind))}: {_number((detail or {}).get('case_count', 0))} cases"
        for kind, detail in counts.items()
    ) or "None"
    return [
        ["Variation coverage", _percent(item.get("variation_coverage"))],
        ["Variation cases", variation_summary],
        ["Accuracy", _percent(item.get("accuracy"))],
        ["Consistency", _percent(item.get("consistency"))],
        ["Worst confidence drop", _percent(item.get("worst_confidence_drop"))],
        ["Missing variation types", ", ".join(item.get("missing_variation_types") or []) or "None"],
        ["Failed case IDs", ", ".join(item.get("failed_case_ids") or []) or "None"],
    ]


def _cost_efficiency_rows(metrics: dict) -> list[list[str]]:
    item = metrics.get("cost_efficiency") or {}
    if item.get("measurement_status") != "measured":
        return [["Status", "Not measurable"], ["Reason", _safe_text(item.get("reason"))]]
    return [
        ["Cost source", _label(str(item.get("cost_source")))],
        ["Measured cases", _number(item.get("case_count"))],
        ["Total cost", _usd(item.get("total_cost_usd"))],
        ["Cost per case", _usd(item.get("cost_per_case_usd"))],
        ["Cost per correct case", _usd(item.get("cost_per_correct_case_usd"))],
        ["Input / output / total tokens", " / ".join(_number(item.get(key)) for key in ("input_tokens", "output_tokens", "total_tokens"))],
        ["Requests per case", _number(item.get("requests_per_case"))],
        ["Retries / tool calls", " / ".join(_number(item.get(key)) for key in ("retry_count", "tool_call_count"))],
        ["Cache hit rate", _percent(item.get("cache_hit_rate"))],
        ["Fallback rate", _percent(item.get("fallback_rate"))],
        ["Timeout rate", _percent(item.get("timeout_rate"))],
        ["Median / P95 / max latency", " / ".join(f"{_number(item.get(key))} ms" for key in ("median_latency_ms", "p95_latency_ms", "max_latency_ms"))],
    ]


def _gate_rows(metrics: dict) -> list[list[str]]:
    gates = metrics.get("gate_details") or []
    if not gates:
        return [["No gates", "-", "-", "-", "Not measurable"]]
    return [
        [
            _label(str(gate.get("dimension") or "-")),
            _label(str(gate.get("name") or "-")),
            _number(gate.get("actual")),
            f"{gate.get('operator', '')} {_number(gate.get('threshold'))}".strip(),
            _status(gate.get("passed")),
        ]
        for gate in gates
    ]


def _role_rows(metrics: dict) -> list[list[str]]:
    roles = metrics.get("role_results") or {}
    return [
        [
            _label(role_id),
            _safe_text(role.get("implementation")),
            _safe_text(role.get("measurement_status")).replace("_", " ").upper(),
            _safe_text(role.get("verdict")).upper(),
        ]
        for role_id, role in roles.items()
    ] or [["No roles", "-", "-", "-"]]


def _count_label(value: int, singular: str, plural: str | None = None) -> str:
    return f"{value} {singular if value == 1 else (plural or singular + 's')}"


def _process_rows(evaluation: dict, metrics: dict) -> list[list[str]]:
    """Describe the recorded deterministic evaluation procedure, not model reasoning."""
    manifest = evaluation.get("input_manifest") or {}
    expected = manifest.get("expected_labels") or []
    confidences = manifest.get("confidences") or []
    claims = manifest.get("claims") or []
    security_cases = ((manifest.get("security") or {}).get("cases") or [])
    trajectory = manifest.get("trajectory") or {}
    trace = manifest.get("trace_envelope") or {}
    adapter = manifest.get("adapter_provenance") or {}
    client_runner = manifest.get("client_provenance") or {}
    required_dimensions = metrics.get("required_dimensions") or []
    optional_dimensions = ("rag", "robustness", "judge_agreement", "reproducibility", "cost_efficiency")
    gates = metrics.get("gate_details") or []
    passed = sum(gate.get("passed") is True for gate in gates)
    failed = sum(gate.get("passed") is False for gate in gates)
    unavailable = sum(gate.get("passed") is None for gate in gates)
    positive_controls = sum(bool(case.get("expected_detection")) for case in security_cases)
    negative_controls = len(security_cases) - positive_controls
    evidence_references = sum(len(claim.get("evidence_ids") or []) for claim in claims)
    optional_status = ", ".join(
        f"{_label(dimension)}: {_safe_text((metrics.get(dimension) or {}).get('measurement_status', 'not_supplied')).replace('_', ' ')}"
        for dimension in optional_dimensions
    )
    return [
        [
            "1. Validate and pin the request",
            f"Validated evaluator schema {metrics.get('schema_version', '-')} and pinned the workspace project, subject, dataset, policy, and evaluator versions in the immutable evaluation snapshot.",
        ],
        [
            "2. Verify integration provenance",
            f"Recorded subject type {_safe_text(manifest.get('subject_type', 'agent'))} through {_safe_text(manifest.get('integration_mode', 'manifest'))}. "
            + (
                f"Invoked approved adapter {_safe_text(adapter.get('adapter_name'))} for {_count_label(adapter.get('case_count', 0), 'bounded case')}; retained request digest {_safe_text(adapter.get('request_sha256'), 80)}, response digest {_safe_text(adapter.get('response_sha256'), 80)}, and duration {_safe_text(adapter.get('duration_ms'))} ms."
                if adapter else
                f"Verified client-runner package {_safe_text(client_runner.get('package_id'))} from {_safe_text(client_runner.get('identity_type'))} identity {_safe_text(client_runner.get('identity_name'))}; retained package digest {_safe_text(client_runner.get('package_sha256'), 80)} and execution digest {_safe_text(client_runner.get('execution_sha256'), 80)}."
                if client_runner else
                f"Verified {_count_label(len(trace.get('events') or []), 'redacted trace event')} against trace digest {_safe_text(trace.get('content_sha256'), 80)}."
                if trace else "No agent trace was supplied for this evaluation."
            ),
        ],
        [
            "3. Classify and calibrate",
            f"Compared {_count_label(len(expected), 'labelled expected/predicted pair')} and {_count_label(len(confidences), 'confidence value')}; classification and calibration are {_safe_text((metrics.get('classification') or {}).get('measurement_status', 'not_measurable')).replace('_', ' ')}.",
        ],
        [
            "4. Check evidence grounding",
            f"Reviewed {_count_label(len(claims), 'declared claim')} with {_count_label(evidence_references, 'evidence reference')}, including citation and artifact-integrity checks.",
        ],
        [
            "5. Test security verdicts",
            f"Evaluated {_count_label(len(security_cases), 'labelled security case')}: {_count_label(positive_controls, 'expected-detection control')} and {_count_label(negative_controls, 'negative control')}.",
        ],
        [
            "6. Inspect trajectory policy",
            f"Compared {_count_label(len(trajectory.get('observed_milestones') or []), 'observed milestone')} with {_count_label(len(trajectory.get('required_milestones') or []), 'required milestone')}; recorded {_count_label(trajectory.get('action_count', 0), 'action')}.",
        ],
        [
            "7. Record optional coverage",
            optional_status,
        ],
        [
            "8. Apply release policy",
            f"Applied {len(gates)} required gate(s) across {_safe_text(', '.join(_label(dimension) for dimension in required_dimensions) or 'no dimensions')}: {passed} passed, {failed} failed, {unavailable} unavailable. The resulting decision is {_safe_text(metrics.get('release_decision', evaluation.get('release_decision'))).upper()}.",
        ],
    ]


def _provenance_rows(evaluation: dict) -> list[list[str]]:
    """Render public governance metadata without exposing dataset or trace content."""
    dataset = evaluation.get("dataset") or {}
    manifest = evaluation.get("input_manifest") or {}
    trace = manifest.get("trace_envelope") or {}
    adapter = manifest.get("adapter_provenance") or {}
    client_runner = manifest.get("client_provenance") or {}
    rows = [
        ["Evaluation project", _safe_text(evaluation.get("project_key", "default"))],
        ["Dataset name", _safe_text(dataset.get("name", "Historical dataset"))],
        ["Dataset classification", _safe_text(dataset.get("classification", "Not recorded"))],
        ["Dataset source reference", _safe_text(dataset.get("source_reference", "Not recorded"))],
        ["Dataset source SHA-256", _safe_text(dataset.get("source_sha256", "Not recorded"))],
        ["Dataset approval status", _safe_text(dataset.get("status", "Historical record"))],
        ["Subject type", _safe_text(manifest.get("subject_type", "agent"))],
        ["Integration mode", _safe_text(manifest.get("integration_mode", "manifest"))],
    ]
    if trace:
        rows.extend([
            ["Trace schema", _safe_text(trace.get("schema_version"))],
            ["Trace producer", _safe_text(trace.get("producer"))],
            ["Redacted trace event count", _number(trace.get("event_count"))],
            ["Trace SHA-256", _safe_text(trace.get("content_sha256"))],
        ])
    if adapter:
        rows.extend([
            ["Live adapter", _safe_text(adapter.get("adapter_name"))],
            ["Adapter contract", _safe_text(adapter.get("adapter_type"))],
            ["Adapter endpoint SHA-256", _safe_text(adapter.get("endpoint_sha256"))],
            ["Adapter request SHA-256", _safe_text(adapter.get("request_sha256"))],
            ["Adapter response SHA-256", _safe_text(adapter.get("response_sha256"))],
            ["Bounded cases / duration", f"{_number(adapter.get('case_count'))} / {_number(adapter.get('duration_ms'))} ms"],
            ["Egress network policy", _safe_text(adapter.get("network_policy_version"))],
        ])
    if client_runner:
        rows.extend([
            ["Client runner identity type", _safe_text(client_runner.get("identity_type"))],
            ["Client runner identity", _safe_text(client_runner.get("identity_name"))],
            ["Client identity fingerprint", _safe_text(client_runner.get("identity_fingerprint"))],
            ["Client package ID", _safe_text(client_runner.get("package_id"))],
            ["Client package SHA-256", _safe_text(client_runner.get("package_sha256"))],
            ["Client execution SHA-256", _safe_text(client_runner.get("execution_sha256"))],
            ["Client source attestation SHA-256", _safe_text(client_runner.get("source_attestation_sha256"))],
            ["Client runner version", _safe_text(client_runner.get("runner_version"))],
        ])
    return rows


def report_context(evaluation: dict) -> dict:
    metrics = dict(evaluation.get("metrics") or {})
    return {
        "evaluation": evaluation,
        "metrics": metrics,
        "source_sha256": evaluation_source_sha256(evaluation),
        "input_manifest_sha256": _input_manifest_sha256(evaluation),
        "metrics_sha256": _metrics_sha256(evaluation),
    }


def _set_cell_fill(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_border(cell) -> None:
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    cell._tc.get_or_add_tcPr().append(borders)


def _set_cell_text(cell, value: object, *, bold: bool = False, color: str = "000000") -> None:
    cell.text = _safe_text(value, 3000)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_border(cell)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(2)
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Aptos"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(color)


def _doc_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_fill(cell, "16324F")
        _set_cell_text(cell, header, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                _set_cell_fill(cells[index], "F4F7F9")
            _set_cell_text(cells[index], value)


def _doc_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    run.font.size = Pt(15 if level == 1 else 11)
    run.font.color.rgb = RGBColor(0, 0, 0)


def _doc_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(_safe_text(text, 5000))
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(10.5)


def _doc_page_number(document: Document) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("ExposureScopeX | Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def render_evaluation_docx(context: dict) -> bytes:
    evaluation, metrics = context["evaluation"], context["metrics"]
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.72)
    _doc_page_number(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI Assurance Pre-release Report")
    run.font.name = "Aptos Display"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("ExposureScopeX evidence-backed AI release decision").italic = True
    for item in subtitle.runs:
        item.font.name = "Aptos"
        item.font.size = Pt(11)
    document.add_paragraph()
    _doc_table(document, ["Evaluated subject", "Version", "Dataset", "Decision"], [[
        evaluation.get("evaluated_agent_id"),
        evaluation.get("evaluated_agent_version"),
        evaluation.get("dataset_version"),
        _safe_text(evaluation.get("release_decision")).upper(),
    ]])
    _doc_body(document, "This report records the deterministic release decision for the stated evaluated subject, version, labelled dataset, policy, and evaluator version. It does not control assessment scanners or make claims beyond the submitted evaluation data.")
    document.add_page_break()

    _doc_heading(document, "Contents")
    _doc_table(document, ["Order", "Section"], [[str(index), title] for index, title in enumerate((
        "Document Control", "Dataset and Integration Provenance", "Evaluation Process and Decision Trace", "Executive Decision", "Release Gate Results",
        "Classification Quality", "Calibration and Evidence Grounding", "Security Outcome and Trajectory",
        "RAG, Robustness and Efficiency", "Optional Measurements", "Limitations and Integrity Manifest",
    ), start=1)])

    _doc_heading(document, "Document Control")
    _doc_table(document, ["Field", "Recorded value"], [
        ["Evaluation name", evaluation.get("name")],
        ["Evaluation ID", evaluation.get("id")],
        ["Evaluated subject", evaluation.get("evaluated_agent_id")],
        ["Subject version", evaluation.get("evaluated_agent_version")],
        ["Evaluator", evaluation.get("evaluator_agent_id")],
        ["Evaluator version", evaluation.get("evaluator_version")],
        ["Dataset version", evaluation.get("dataset_version")],
        ["Evaluation timestamp", _format_timestamp(evaluation.get("created_at"))],
        ["Report renderer version", EVALUATOR_REPORT_VERSION],
    ])

    _doc_heading(document, "Dataset and Integration Provenance")
    _doc_body(document, "This section records the governed source identity and integration metadata. It deliberately excludes raw dataset rows, prompts, responses, and trace content.")
    _doc_table(document, ["Control", "Recorded value"], _provenance_rows(evaluation))

    _doc_heading(document, "Evaluation Process and Decision Trace")
    _doc_body(document, "This auditable trace is derived from the stored request and calculated metric snapshot. It identifies the inputs and deterministic checks that produced the decision; it is not hidden model reasoning or a claim of unrecorded execution.")
    _doc_table(document, ["Step", "Recorded procedure"], _process_rows(evaluation, metrics))

    _doc_heading(document, "Executive Decision")
    decision = _safe_text(evaluation.get("release_decision")).upper()
    if decision == "PASS":
        _doc_body(document, "The evaluated subject met every required, measurable gate in the declared evaluation policy for this specific labelled dataset. This decision is evidence of conformance to the declared benchmark only; it is not a general safety, security, or production guarantee.")
    elif decision == "FAIL":
        _doc_body(document, "The evaluated subject did not meet one or more required measurable release gates. Release should remain blocked until the failed cases, thresholds, and corrective changes have been reviewed and a new version is evaluated.")
    else:
        _doc_body(document, "The evaluation is inconclusive because at least one required dimension could not be measured. Missing ground truth or evidence never produces a passing release decision.")
    _doc_table(document, ["Measure", "Result"], _metric_rows(metrics))

    _doc_heading(document, "Release Gate Results")
    _doc_table(document, ["Dimension", "Gate", "Actual", "Threshold", "Result"], _gate_rows(metrics))
    _doc_heading(document, "Logical Role Results", 2)
    _doc_table(document, ["Role", "Implementation", "Measurement", "Verdict"], _role_rows(metrics))

    _doc_heading(document, "Classification Quality")
    _doc_table(document, ["Measure", "Result"], _classification_rows(metrics))
    _doc_heading(document, "Confusion Matrix", 2)
    _doc_table(document, ["Expected label", "Predicted label", "Count"], _confusion_rows(metrics))

    _doc_heading(document, "Calibration and Evidence Grounding")
    _doc_heading(document, "Confidence Calibration", 2)
    _doc_table(document, ["Measure", "Result"], _confidence_rows(metrics))
    _doc_heading(document, "Evidence Grounding", 2)
    _doc_table(document, ["Measure", "Result"], _grounding_rows(metrics))

    _doc_heading(document, "Security Outcome and Trajectory")
    _doc_heading(document, "Security Verdict", 2)
    _doc_table(document, ["Measure", "Result"], _security_rows(metrics))
    _doc_heading(document, "Trajectory Policy", 2)
    _doc_table(document, ["Measure", "Result"], _trajectory_rows(metrics))

    _doc_heading(document, "RAG, Robustness and Efficiency")
    _doc_heading(document, "RAG Quality", 2)
    _doc_table(document, ["Measure", "Result"], _rag_rows(metrics))
    _doc_heading(document, "Robustness Coverage", 2)
    _doc_table(document, ["Measure", "Result"], _robustness_rows(metrics))
    _doc_heading(document, "Cost and Efficiency", 2)
    _doc_table(document, ["Measure", "Result"], _cost_efficiency_rows(metrics))

    _doc_heading(document, "Optional Measurements")
    _doc_table(document, ["Dimension", "Status", "Recorded result or limitation"], _optional_dimension_rows(metrics))

    _doc_heading(document, "Limitations and Integrity Manifest")
    for limitation in metrics.get("limitations") or ["No limitations were recorded."]:
        _doc_body(document, f"- {_safe_text(limitation)}")
    _doc_table(document, ["Element", "SHA-256 or identifier"], [
        ["Immutable evaluation snapshot", context["source_sha256"]],
        ["Submitted input manifest", context["input_manifest_sha256"]],
        ["Calculated metrics", context["metrics_sha256"]],
        ["Evaluation ID", evaluation.get("id")],
    ])
    _doc_body(document, "The report is generated from the stored evaluation snapshot. The download endpoint verifies the report file digest before delivery. The original evaluation record remains the authoritative source for all metric values.")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_text(value: object, limit: int = 3000) -> str:
    return escape(_safe_text(value, limit)).replace("\n", "<br/>")


def _pdf_paragraph(value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(_pdf_text(value), style)


def _pdf_table(headers: list[str], rows: list[list[str]], body: ParagraphStyle, *, widths: list[float] | None = None) -> Table:
    rendered = [[_pdf_paragraph(header, body) for header in headers]]
    rendered.extend([[_pdf_paragraph(value, body) for value in row] for row in rows])
    table = Table(rendered, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands: list[tuple] = [
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for index in range(1, len(rendered)):
        if index % 2 == 0:
            commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#F4F7F9")))
    table.setStyle(TableStyle(commands))
    return table


def _pdf_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#4B5E6B"))
    canvas.drawString(42, 25, "ExposureScopeX | AI Assurance Pre-release Report")
    canvas.drawRightString(letter[0] - 42, 25, f"Page {document.page}")
    canvas.restoreState()


def render_evaluation_pdf(context: dict) -> bytes:
    evaluation, metrics = context["evaluation"], context["metrics"]
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=50, bottomMargin=45
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("EvaluatorTitle", parent=styles["Title"], fontName="Helvetica", fontSize=22, leading=27, textColor=colors.black, alignment=1, spaceAfter=8)
    subtitle = ParagraphStyle("EvaluatorSubtitle", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=14, textColor=colors.HexColor("#344955"), alignment=1, spaceAfter=18)
    heading = ParagraphStyle("EvaluatorHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.black, spaceBefore=14, spaceAfter=7)
    subheading = ParagraphStyle("EvaluatorSubheading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.black, spaceBefore=10, spaceAfter=5)
    body = ParagraphStyle("EvaluatorBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, textColor=colors.black, spaceAfter=6)
    table_body = ParagraphStyle("EvaluatorTable", parent=body, fontSize=7.5, leading=9)
    story = [
        Paragraph("AI Assurance Pre-release Report", title),
        Paragraph("ExposureScopeX evidence-backed AI release decision", subtitle),
        _pdf_table(["Evaluated subject", "Version", "Dataset", "Decision"], [[
            _safe_text(evaluation.get("evaluated_agent_id")),
            _safe_text(evaluation.get("evaluated_agent_version")),
            _safe_text(evaluation.get("dataset_version")),
            _safe_text(evaluation.get("release_decision")).upper(),
        ]], table_body, widths=[1.65 * inch, 1.25 * inch, 2.45 * inch, 1.15 * inch]),
        Spacer(1, 12),
        Paragraph("This report records the deterministic release decision for the stated evaluated subject, version, labelled dataset, policy, and evaluator version. It does not control assessment scanners or make claims beyond the submitted evaluation data.", body),
        PageBreak(),
        Paragraph("Document Control", heading),
        _pdf_table(["Field", "Recorded value"], [
            ["Evaluation name", _safe_text(evaluation.get("name"))],
            ["Evaluation ID", _safe_text(evaluation.get("id"))],
            ["Evaluated subject", _safe_text(evaluation.get("evaluated_agent_id"))],
            ["Subject version", _safe_text(evaluation.get("evaluated_agent_version"))],
            ["Evaluator", _safe_text(evaluation.get("evaluator_agent_id"))],
            ["Evaluator version", _safe_text(evaluation.get("evaluator_version"))],
            ["Dataset version", _safe_text(evaluation.get("dataset_version"))],
            ["Evaluation timestamp", _format_timestamp(evaluation.get("created_at"))],
            ["Report renderer version", EVALUATOR_REPORT_VERSION],
        ], table_body, widths=[2.05 * inch, 5.0 * inch]),
        Paragraph("Dataset and Integration Provenance", heading),
        Paragraph("This section records the governed source identity and integration metadata. It deliberately excludes raw dataset rows, prompts, responses, and trace content.", body),
        _pdf_table(["Control", "Recorded value"], _provenance_rows(evaluation), table_body, widths=[2.05 * inch, 5.0 * inch]),
        Paragraph("Evaluation Process and Decision Trace", heading),
        Paragraph("This auditable trace is derived from the stored request and calculated metric snapshot. It identifies the inputs and deterministic checks that produced the decision; it is not hidden model reasoning or a claim of unrecorded execution.", body),
        _pdf_table(["Step", "Recorded procedure"], _process_rows(evaluation, metrics), table_body, widths=[2.05 * inch, 5.0 * inch]),
        Paragraph("Executive Decision", heading),
        Paragraph(
            "The evaluated subject met every required, measurable gate in the declared policy for this specific labelled dataset. This decision is benchmark-specific and is not a general safety, security, or production guarantee."
            if evaluation.get("release_decision") == "pass"
            else "The evaluation did not produce a passing release decision. Review failed or unavailable required dimensions before using this result as a release gate.",
            body,
        ),
        _pdf_table(["Measure", "Result"], _metric_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        Paragraph("Release Gate Results", heading),
        _pdf_table(["Dimension", "Gate", "Actual", "Threshold", "Result"], _gate_rows(metrics), table_body, widths=[1.2 * inch, 1.75 * inch, 1.0 * inch, 1.9 * inch, 1.2 * inch]),
        Paragraph("Logical Role Results", subheading),
        _pdf_table(["Role", "Implementation", "Measurement", "Verdict"], _role_rows(metrics), table_body, widths=[1.55 * inch, 2.05 * inch, 1.85 * inch, 1.6 * inch]),
        Paragraph("Classification Quality", heading),
        _pdf_table(["Measure", "Result"], _classification_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        KeepTogether([
            Paragraph("Confusion Matrix", subheading),
            _pdf_table(["Expected label", "Predicted label", "Count"], _confusion_rows(metrics), table_body, widths=[2.6 * inch, 2.6 * inch, 1.85 * inch]),
        ]),
        Paragraph("Calibration and Evidence Grounding", heading),
        Paragraph("Confidence Calibration", subheading),
        _pdf_table(["Measure", "Result"], _confidence_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        # Keep the evidence exhibit label with its compact table. A section heading
        # without its exhibit is ambiguous when a report spans multiple pages.
        KeepTogether([
            Paragraph("Evidence Grounding", subheading),
            _pdf_table(["Measure", "Result"], _grounding_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        ]),
        Paragraph("Security Outcome and Trajectory", heading),
        Paragraph("Security Verdict", subheading),
        _pdf_table(["Measure", "Result"], _security_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        KeepTogether([
            Paragraph("Trajectory Policy", subheading),
            _pdf_table(["Measure", "Result"], _trajectory_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        ]),
        Paragraph("RAG, Robustness and Efficiency", heading),
        Paragraph("RAG Quality", subheading),
        _pdf_table(["Measure", "Result"], _rag_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        KeepTogether([
            Paragraph("Robustness Coverage", subheading),
            _pdf_table(["Measure", "Result"], _robustness_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        ]),
        KeepTogether([
            Paragraph("Cost and Efficiency", subheading),
            _pdf_table(["Measure", "Result"], _cost_efficiency_rows(metrics), table_body, widths=[2.5 * inch, 4.55 * inch]),
        ]),
        Paragraph("Optional Measurements", heading),
        _pdf_table(["Dimension", "Status", "Recorded result or limitation"], _optional_dimension_rows(metrics), table_body, widths=[1.35 * inch, 1.3 * inch, 4.4 * inch]),
        Paragraph("Limitations and Integrity Manifest", heading),
    ]
    for limitation in metrics.get("limitations") or ["No limitations were recorded."]:
        story.append(Paragraph(f"- {_pdf_text(limitation)}", body))
    story.append(_pdf_table(["Element", "SHA-256 or identifier"], [
        ["Immutable evaluation snapshot", context["source_sha256"]],
        ["Submitted input manifest", context["input_manifest_sha256"]],
        ["Calculated metrics", context["metrics_sha256"]],
        ["Evaluation ID", _safe_text(evaluation.get("id"))],
    ], table_body, widths=[2.05 * inch, 5.0 * inch]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("The report is generated from the stored evaluation snapshot. The download endpoint verifies the report file digest before delivery. The original evaluation record remains the authoritative source for all metric values.", body))
    document.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return buffer.getvalue()


def _safe_filename(value: object) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", _safe_text(value, 100)).strip("-") or "evaluation"


def evaluation_report_path(storage_key: str) -> Path:
    from .config import settings

    root = (Path(settings().artifact_root) / "evaluator-reports").resolve()
    candidate = (root / storage_key).resolve()
    if not candidate.is_relative_to(root):
        raise EvaluationReportError("Evaluator report path escapes configured storage root")
    return candidate


async def get_evaluation_report(evaluation_id: UUID, report_format: ReportFormat) -> dict | None:
    """Return the immutable report record for the active renderer release."""
    from .db import pool

    row = await pool().fetchrow(
        """
        SELECT id, evaluation_id, report_format, renderer_version, source_sha256,
               content_sha256, storage_key, size_bytes, generated_at
        FROM evaluation_reports
        WHERE evaluation_id = $1 AND report_format = $2 AND renderer_version = $3
        """,
        evaluation_id, report_format, EVALUATOR_REPORT_VERSION,
    )
    return dict(row) if row else None


async def list_evaluation_reports(evaluation_id: UUID) -> list[dict]:
    from .db import pool

    rows = await pool().fetch(
        """
        SELECT id, evaluation_id, report_format, renderer_version, source_sha256,
               content_sha256, storage_key, size_bytes, generated_at
        FROM evaluation_reports
        WHERE evaluation_id = $1 AND renderer_version = $2
        ORDER BY report_format
        """,
        evaluation_id, EVALUATOR_REPORT_VERSION,
    )
    return [dict(row) for row in rows]


async def _persist_report(
    evaluation: dict,
    report_format: ReportFormat,
    content: bytes,
    source_sha256: str,
    generated_by: UUID,
) -> dict:
    from .db import pool

    evaluation_id = evaluation["id"]
    if existing := await get_evaluation_report(evaluation_id, report_format):
        return existing
    content_sha256 = hashlib.sha256(content).hexdigest()
    safe_subject = _safe_filename(evaluation.get("evaluated_agent_id"))
    renderer_tag = EVALUATOR_REPORT_VERSION.replace(".", "-")
    filename = f"{safe_subject}-{evaluation_id}-r{renderer_tag}-{source_sha256[:12]}.{REPORT_EXTENSIONS[report_format]}"
    storage_key = f"{evaluation_id}/{filename}"
    destination = evaluation_report_path(storage_key)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not destination.exists():
        temporary = destination.with_suffix(destination.suffix + ".tmp")
        temporary.write_bytes(content)
        os.replace(temporary, destination)
    inserted = await pool().fetchrow(
        """
        INSERT INTO evaluation_reports (
          evaluation_id, report_format, renderer_version, source_sha256,
          content_sha256, storage_key, size_bytes, generated_by
        ) VALUES ($1, $2, $3, $4, $5, $6, $7, $8)
        ON CONFLICT (evaluation_id, report_format, renderer_version) DO NOTHING
        RETURNING id, evaluation_id, report_format, renderer_version, source_sha256,
                  content_sha256, storage_key, size_bytes, generated_at
        """,
        evaluation_id, report_format, EVALUATOR_REPORT_VERSION, source_sha256,
        content_sha256, storage_key, len(content), generated_by,
    )
    if inserted:
        return dict(inserted)
    existing = await get_evaluation_report(evaluation_id, report_format)
    if existing is None:
        raise EvaluationReportError("Evaluator report record could not be persisted")
    return existing


async def generate_evaluation_reports(evaluation: dict, generated_by: UUID) -> dict[str, dict]:
    """Generate each immutable report format once for a stored evaluation."""
    evaluation_id = evaluation.get("id")
    if not isinstance(evaluation_id, UUID):
        raise EvaluationReportError("Stored evaluation is missing a UUID identifier")
    context = report_context(evaluation)
    source_sha256 = context["source_sha256"]
    pending = [
        report_format
        for report_format in ("docx", "pdf")
        if await get_evaluation_report(evaluation_id, report_format) is None
    ]
    generated: dict[str, bytes] = {}
    if "docx" in pending:
        generated["docx"] = await asyncio.to_thread(render_evaluation_docx, context)
    if "pdf" in pending:
        generated["pdf"] = await asyncio.to_thread(render_evaluation_pdf, context)
    records: dict[str, dict] = {}
    for report_format in ("docx", "pdf"):
        records[report_format] = (
            await _persist_report(
                evaluation, report_format, generated[report_format], source_sha256, generated_by
            )
            if report_format in generated
            else await get_evaluation_report(evaluation_id, report_format)
        )
    return records
