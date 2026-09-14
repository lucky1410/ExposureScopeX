from __future__ import annotations

import asyncio
import hashlib
import io
import json
import zipfile
from html import escape
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from PIL import Image as PILImage
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

from .artifact_store import artifact_path, persist_bytes
from .benchmark_manifest import benchmark_release_eligibility, load_benchmark_manifest, map_observation_key
from .benchmarking import BenchmarkObservation, score_vulnerability_benchmark
from .config import settings
from .db import pool
from .methodology_cases import CASES_BY_ADAPTER, canonical_observation_key, observation_evidence_kinds
from .redaction import redact_object, redact_text
from .serialization import json_safe
from .target_planning import canonical_http_url


SEVERITIES = ("critical", "high", "medium", "low", "info")
SEVERITY_COLORS = {"critical": "8B1E2D", "high": "C2413B", "medium": "C97A16", "low": "2F6F8F", "info": "607481"}
REPORT_TERMINAL_STATUSES = {
    "complete": "final",
    "partial": "partial",
    "failed": "failed",
    "cancelled": "cancelled",
    "blocked": "blocked",
}
REPORT_SECTIONS = [
    "Document Control",
    "Executive Summary",
    "Project Scope",
    "Profiling",
    "Finding Index",
    "Detailed Findings",
    "Critical Risk Findings",
    "High Risk Findings",
    "Medium Risk Findings",
    "Low Risk Findings",
    "Informational Findings",
    "Benchmark Evaluation",
    "Appendix A: Engagement Methodology",
    "Appendix B: Risk Methodology",
    "Appendix C: Testing Methodologies",
    "Appendix D: Execution Coverage and Exceptions",
    "Appendix E: Evidence Integrity Manifest",
]

FINDING_GUIDANCE = {
    "http.header.absent:strict-transport-security": {
        "impact": "Without HTTP Strict Transport Security, a browser that first reaches the site over HTTP may be more exposed to protocol-downgrade or interception risks. Its practical impact depends on HTTP reachability, redirect behavior, and client deployment context.",
        "cvss": "4.3", "vector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:R/S:U/C:L/I:N/A:N",
        "cwe": "CWE-319: Cleartext Transmission of Sensitive Information", "owasp": "OWASP Top 10 2021 A05: Security Misconfiguration",
        "references": ["https://developer.mozilla.org/docs/Web/HTTP/Headers/Strict-Transport-Security", "https://cwe.mitre.org/data/definitions/319.html"],
    },
    "http.header.absent:content-security-policy": {
        "impact": "Without Content-Security-Policy, the browser has no application-defined restriction on permitted script, frame, style, and connection sources. This does not create cross-site scripting by itself, but it removes an important compensating control and can increase the impact of a separate injection weakness.",
        "cvss": "4.3", "vector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:R/S:U/C:L/I:N/A:N",
        "cwe": "CWE-693: Protection Mechanism Failure", "owasp": "OWASP Top 10 2021 A05: Security Misconfiguration",
        "references": ["https://developer.mozilla.org/docs/Web/HTTP/CSP", "https://cwe.mitre.org/data/definitions/693.html"],
    },
    "http.header.absent:x-content-type-options": {
        "impact": "Browsers may MIME-sniff a response instead of honoring its declared media type. Where attacker-influenced content is served, this can cause data to be interpreted as active content. Exploitability depends on upload, reflection, and content-serving behavior.",
        "cvss": "3.1", "vector": "CVSS:3.1/AV:N/AC:H/PR:N/UI:R/S:U/C:L/I:N/A:N",
        "cwe": "CWE-16: Configuration", "owasp": "OWASP Top 10 2021 A05: Security Misconfiguration",
        "references": ["https://developer.mozilla.org/docs/Web/HTTP/Headers/X-Content-Type-Options", "https://cwe.mitre.org/data/definitions/16.html"],
    },
    "http.header.absent:x-frame-options": {
        "impact": "The assessed response can be framed unless an effective Content-Security-Policy frame-ancestors directive provides equivalent protection. A hostile site may overlay or disguise the framed interface and induce an authenticated user to perform unintended actions.",
        "cvss": "4.3", "vector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:R/S:U/C:N/I:L/A:N",
        "cwe": "CWE-1021: Improper Restriction of Rendered UI Layers or Frames", "owasp": "OWASP Top 10 2021 A05: Security Misconfiguration",
        "references": ["https://developer.mozilla.org/docs/Web/HTTP/Headers/X-Frame-Options", "https://cwe.mitre.org/data/definitions/1021.html"],
    },
    "http.header.absent:referrer-policy": {
        "impact": "Without an explicit Referrer-Policy, browsers apply their default policy and may disclose origin or path information during outbound navigation. Sensitive data placed in URL paths or queries can therefore be exposed to unintended destinations, subject to browser behavior.",
        "cvss": "3.1", "vector": "CVSS:3.1/AV:N/AC:H/PR:N/UI:R/S:U/C:L/I:N/A:N",
        "cwe": "CWE-200: Exposure of Sensitive Information", "owasp": "OWASP Top 10 2021 A05: Security Misconfiguration",
        "references": ["https://developer.mozilla.org/docs/Web/HTTP/Headers/Referrer-Policy", "https://cwe.mitre.org/data/definitions/200.html"],
    },
    "nuclei.template:phpinfo-files": {
        "impact": "A public PHPInfo page discloses runtime versions, modules, filesystem paths, environment characteristics, and security-relevant configuration. This materially improves attacker reconnaissance and may expose secrets when environment variables or request data are displayed.",
        "cvss": "5.3", "vector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:L/I:N/A:N",
        "cwe": "CWE-200: Exposure of Sensitive Information", "owasp": "OWASP Top 10 2021 A05: Security Misconfiguration",
        "references": ["https://www.php.net/manual/en/function.phpinfo.php", "https://cwe.mitre.org/data/definitions/200.html"],
    },
}


def _safe_text(value, limit: int = 4000) -> str:
    text = (str(value or "-").replace("\x00", "").replace("·", "|")
            .replace("–", "-").replace("—", "-").replace("‑", "-"))
    return text[:limit]


def _display_url(value: object) -> str:
    text = str(value or "-")
    return canonical_http_url(text) if text.startswith(("http://", "https://")) else text


def report_status_for_scan(scan_status: str) -> str:
    mapped = REPORT_TERMINAL_STATUSES.get(str(scan_status))
    if mapped is None:
        raise ValueError(f"unsupported terminal scan status for report generation: {scan_status}")
    return mapped


def _format_timestamp(value: object) -> str:
    text = str(value or "-")
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    except ValueError:
        return _safe_text(text)


def _finding_key(finding: dict) -> str | None:
    evidence = finding.get("evidence") or {}
    return finding.get("canonical_observation_key") or canonical_observation_key(evidence)


def _finding_guidance(finding: dict) -> dict:
    return FINDING_GUIDANCE.get(_finding_key(finding), {})


def _finding_business_impact(finding: dict) -> str:
    return _finding_guidance(finding).get("impact") or finding.get("business_impact") or "Business impact requires client asset and control context."


def _finding_taxonomy(finding: dict) -> list[list[str]]:
    guidance = _finding_guidance(finding)
    return [[
        guidance.get("cvss", "Not assigned"), guidance.get("vector", "Not assigned"),
        guidance.get("cwe", "Not mapped"), guidance.get("owasp", "Not mapped"),
    ]]


def _finding_references(finding: dict) -> list[str]:
    references = list((finding.get("evidence") or {}).get("references") or [])
    references.extend(_finding_guidance(finding).get("references") or [])
    return list(dict.fromkeys(str(item) for item in references if item))


def _profile_methodology_text(profile: str) -> str:
    common = (
        "exact-origin validation; HTTP response profiling; profile-bounded TCP service and TLS discovery; "
        "same-origin browser crawling with destructive routes excluded, using authentication only when authorized credentials are supplied; deterministic security-header review; "
        "execution of a pinned, inventoried, signed, non-intrusive Nuclei template profile; and post-execution evidence validation"
    )
    additions = {
        "light": "; bounded configuration and exposure checks, including a passive certificate-transparency inventory of descendants of the exact authorized hostname only",
        "medium": "; GET-only application route, form, parameter, and client-resource inventory; authenticated session-cookie control review; and published API-contract review",
        "aggressive": "; GET-only application route, form, parameter, and client-resource inventory; authenticated session-cookie control review; published API-contract review; and route-level HTTP policy consistency review across the expanded approved surface",
    }
    bounded = additions.get(profile, "; profile-specific observational checks")
    return (
        f"{profile.capitalize()} coverage includes {common}{bounded}. "
        "The selected template inventory and discovered in-scope URLs define coverage. Runtime is only a watchdog boundary and is not used as a substitute for template completion. "
        "All additional application checks are observational: no forms, payloads, API operations, or state-changing requests are submitted. Manual business-logic testing, exploit chaining, source review, and destructive validation remain outside this profile."
    )


def _testing_methodology_appendix(context: dict) -> str:
    stages = {str(stage["adapter"]) for stage in context["stages"]}
    methods = [
        "Scope preflight verifies the approved exact origin before execution and records the scope manifest.",
        "HTTP profiling performs a bounded GET request and preserves the request, response metadata, and original response artifact.",
        "Service and TLS discovery uses Nmap against the approved host only; it inventories reachable services and TLS posture without authentication attacks or exploitation.",
        "Browser crawl uses Playwright, follows same-origin safe GET routes only, blocks destructive route names, and retains crawl and browser-capture artifacts. It uses an authorized account only when one was configured for the assessment.",
        "Security-header review evaluates observed target response headers. A missing X-Frame-Options finding is suppressed when a CSP frame-ancestors directive provides equivalent protection.",
        "Nuclei runs the pinned, inventoried non-intrusive template set. Its output is normalized into candidates and must satisfy independent evidence validation before being marked confirmed.",
        "Evidence validation verifies artifact hashes and applies deterministic response checks or safe replay. It does not submit forms, payloads, API operations, or state-changing requests.",
    ]
    if "subdomain_enumeration" in stages:
        methods.append("Passive subdomain inventory queries certificate-transparency records for descendants of the exact target hostname. It does not resolve, reach, crawl, or scan discovered names, and an unavailable public source is reported as unavailable rather than a clean inventory.")
    if "application_surface_inventory" in stages:
        methods.append("Application surface inventory records discovered same-origin routes, forms, parameters, and client resources without submitting form data.")
    if "authenticated_session_review" in stages:
        methods.append("Authenticated session review inspects browser-observed cookie attributes only; it does not attempt session takeover or credential attacks.")
    if "api_contract_review" in stages:
        methods.append("API contract review makes bounded GET-only requests to published OpenAPI or Swagger documents and does not invoke application APIs.")
    if "route_security_policy_review" in stages:
        methods.append("Route policy review compares observed HTTP controls across approved crawled routes using GET-only requests.")
    executed = ", ".join(stage["adapter"] for stage in sorted(context["stages"], key=lambda item: item["position"]))
    return f"{_profile_methodology_text(context['assessment']['mode'])}\n\nExecuted stages: {executed}.\n\n" + "\n\n".join(methods)


def _scope_summary(scope: dict | None) -> str:
    if not scope:
        return "No uploaded scope file was supplied; the exact target URL and explicit authorization acknowledgement were the enforced boundary."
    return (
        f"Scope file: {scope['source_format'].upper()} | Authorization: {scope['authorization_id']} | "
        f"Expiry: {scope['authorization_expires_at']} | Allowed ports: {', '.join(str(port) for port in scope['allowed_ports'])} | "
        f"Allowed paths: {', '.join(scope['allowed_paths'])} | Excluded paths: {', '.join(scope['excluded_paths']) or 'None'} | "
        f"Original declaration SHA-256: {scope['source_sha256']}. Credential references are opaque identifiers only; secrets are excluded."
    )


def _finding_index_rows(findings: list[dict]) -> list[list[str]]:
    rows = []
    for number, finding in enumerate(findings, start=1):
        evidence = finding.get("evidence") or {}
        rows.append([
            f"F-{number:03d}",
            str(finding["severity"]).upper(),
            str(finding.get("validation_status") or "candidate").upper(),
            _safe_text(finding["title"], 100),
            _safe_text(_display_url(finding["target"]), 80),
            str(evidence.get("source_sha256") or "-")[:16],
        ])
    return rows or [["-", "-", "-", "No risk findings were recorded.", "-", "-"]]


def _latest_json_artifact(artifacts: list[dict], kind: str) -> dict | None:
    for item in reversed([entry for entry in artifacts if entry.get("kind") == kind]):
        content = _load_artifact_bytes(item)
        if not content:
            continue
        try:
            payload = json.loads(content)
        except (UnicodeDecodeError, json.JSONDecodeError):
            continue
        if isinstance(payload, dict):
            return payload
    return None


def _coverage_family_rows(coverage: list[dict]) -> list[list[object]]:
    families: dict[str, dict[str, int]] = {}
    for record in coverage:
        adapter = str(record.get("adapter") or "")
        methodology_case = CASES_BY_ADAPTER.get(adapter)
        family = str(record.get("family") or (methodology_case.family if methodology_case else "unclassified"))
        bucket = families.setdefault(
            family,
            {"required": 0, "completed": 0, "exceptions": 0, "pending": 0},
        )
        if record.get("required"):
            bucket["required"] += 1
        status = str(record.get("status") or "planned")
        if status in {"completed", "succeeded"}:
            bucket["completed"] += 1
        elif status in {"failed", "timed_out", "skipped", "blocked", "cancelled"}:
            bucket["exceptions"] += 1
        else:
            bucket["pending"] += 1
    rows = [
        [
            _evidence_label(family),
            values["required"],
            values["completed"],
            values["exceptions"],
            values["pending"],
        ]
        for family, values in sorted(families.items())
    ]
    return rows or [["No coverage records", "-", "-", "-", "-"]]


def _coverage_summary_intro(context: dict) -> str:
    profile = str(context["assessment"]["mode"]).capitalize()
    return (
        f"{profile} coverage is measured against methodology families and retained evidence, not the progress bar alone. "
        "Completed means the applicable family produced a terminal successful result. Exceptions include failed, timed out, blocked, skipped, or cancelled required work."
    )


def _profiling_summary_intro(context: dict) -> str:
    profile = str(context["assessment"]["mode"]).capitalize()
    perspective = _assessment_perspective(context)
    return (
        f"{profile} profiling summarizes what the platform actually observed about the approved target from the {perspective.lower()} perspective. "
        "This section is intentionally operational: it distinguishes target reachability, execution quality, methodology-family coverage, and retained profiling artifacts from the vulnerability conclusions presented later in the report."
    )


def _execution_exception_rows(stages: list[dict]) -> list[list[object]]:
    rows = []
    for stage in stages:
        status = str(stage.get("status") or "")
        if status in {"succeeded", "completed"}:
            continue
        rows.append([
            stage.get("position", 0) + 1,
            stage.get("adapter") or "-",
            status or "-",
            stage.get("error_detail") or stage.get("error_code") or "No additional detail recorded",
        ])
    return rows or [["-", "No execution exceptions recorded", "-", "-"]]


def _execution_exception_summary(stages: list[dict]) -> str:
    exceptional = [stage for stage in stages if str(stage.get("status") or "") not in {"succeeded", "completed"}]
    if not exceptional:
        return "All planned stages reached successful terminal states. No execution exceptions were recorded."
    counts: dict[str, int] = {}
    for stage in exceptional:
        status = str(stage.get("status") or "unknown")
        counts[status] = counts.get(status, 0) + 1
    ordered = ", ".join(f"{status}={counts[status]}" for status in sorted(counts))
    return (
        f"Execution exceptions affected {len(exceptional)} planned stage(s). "
        f"Terminal exception counts: {ordered}. "
        "See the exception ledger below for the exact stage-level cause retained by the platform."
    )


def _profile_evidence_sections(context: dict) -> list[dict]:
    artifacts = context["artifacts"]
    sections: list[dict] = []

    standards = _latest_json_artifact(artifacts, "standards_discovery")
    if standards:
        observations = list(standards.get("observations") or [])
        sections.append({
            "title": "Published Metadata Discovery",
            "summary": "Bounded GET-only requests to common security and API metadata locations.",
            "rows": [
                ["Observed locations", sum("status" in item for item in observations)],
                ["Locations with errors", sum("error" in item for item in observations)],
                ["Requested paths", len(observations)],
            ],
        })

    surface = _latest_json_artifact(artifacts, "application_surface_inventory")
    if surface:
        pages = list(surface.get("pages") or [])
        sections.append({
            "title": "Application Surface Inventory",
            "summary": "Same-origin GET-only inventory of application pages, forms, fields, and client resources.",
            "rows": [
                ["Pages inspected", sum("status" in item for item in pages)],
                ["Pages with errors", sum("error" in item for item in pages)],
                ["Forms recorded", sum(len(item.get("forms") or []) for item in pages)],
                ["Fields recorded", sum(len(form.get("fields") or []) for item in pages for form in (item.get("forms") or []))],
                ["External scripts recorded", sum(len(item.get("scripts") or []) for item in pages)],
            ],
        })

    session = _latest_json_artifact(artifacts, "session_review")
    if session:
        authentication = dict(session.get("authentication") or {})
        cookies = list(session.get("cookies") or [])
        sections.append({
            "title": "Authenticated Session Review",
            "summary": "Browser-observed cookie control review using authorized credentials only when supplied.",
            "rows": [
                ["Authentication configured", "Yes" if session.get("authentication_configured") or authentication.get("configured") else "No"],
                ["Authentication verified", "Yes" if authentication.get("verified") else "No"],
                ["Cookies observed", len(cookies)],
                ["Session-like cookies", sum(any(token in str(cookie.get("name") or "").lower() for token in ("session", "sess", "auth", "token", "php")) for cookie in cookies)],
                ["Status", session.get("status") or "completed"],
            ],
        })

    contracts = _latest_json_artifact(artifacts, "api_contract_review")
    if contracts:
        specs = list(contracts.get("contracts") or [])
        schemes = sorted({scheme for item in specs for scheme in (item.get("security_schemes") or [])})
        sections.append({
            "title": "Published API Contract Review",
            "summary": "GET-only discovery of published OpenAPI or Swagger definitions and declared authentication schemes.",
            "rows": [
                ["Conventional locations tested", len(contracts.get("tested_paths") or [])],
                ["Contracts identified", len(specs)],
                ["Declared API paths", sum(int(item.get("path_count") or 0) for item in specs)],
                ["Security schemes", ", ".join(schemes) or "None declared"],
            ],
        })

    route_policy = _latest_json_artifact(artifacts, "route_security_policy_review")
    if route_policy:
        routes = list(route_policy.get("routes") or [])
        sections.append({
            "title": "Route Policy Consistency Review",
            "summary": "GET-only comparison of HTTP security controls across approved crawled routes.",
            "rows": [
                ["Routes reviewed", sum("status" in item for item in routes)],
                ["Routes with errors", sum("error" in item for item in routes)],
                ["Routes missing one or more policy headers", sum(bool(item.get("missing_headers")) for item in routes if "status" in item)],
            ],
        })
    return sections


def _pdf_text(value, limit: int = 4000) -> str:
    return escape(_safe_text(value, limit)).replace("\n", "<br/>")


def _doc_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(16 if level == 1 else 12)
    run.font.color.rgb = RGBColor.from_string("16324F" if level == 1 else "1E5F8A")


def _doc_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(_safe_text(text))
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9)


def _doc_protocol(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(7.5)


def _doc_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "16324F")
        cell._tc.get_or_add_tcPr().append(shading)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = _safe_text(value, 1000)


def _evidence_label(key: str) -> str:
    acronyms = {"api": "API", "http": "HTTP", "id": "ID", "sha256": "SHA-256", "tls": "TLS", "url": "URL"}
    return " ".join(acronyms.get(word, word.title()) for word in key.replace("_", " ").strip().lower().split())


def _evidence_value(value: object, limit: int = 1200) -> str:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list) and all(not isinstance(item, (dict, list)) for item in value):
        return _safe_text(", ".join(str(item) for item in value), limit)
    if isinstance(value, (dict, list)):
        return _safe_text(json.dumps(value, indent=2, default=str), limit)
    return _safe_text(value, limit)


def _protocol_text(value: object, limit: int = 1800) -> str:
    text = redact_text(str(value or "")).replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    if not text:
        return ""
    head, separator, body = text.partition("\n\n")
    if not separator:
        return text[:limit] + (f"\n\n[TRUNCATED FOR READABILITY - complete {len(text)} character original retained in the evidence bundle]" if len(text) > limit else "")
    if not body:
        return head[:limit] + (f"\n\n[TRUNCATED FOR READABILITY - complete {len(head)} character original retained in the evidence bundle]" if len(head) > limit else "")
    # Long encoded lines and embedded assets are forensic payload, not readable
    # report evidence. The exact bytes remain in the hashed evidence bundle.
    body_lines = []
    for line in body.splitlines():
        if "base64," in line.lower():
            body_lines.append("[EMBEDDED BINARY DATA OMITTED]")
        else:
            body_lines.append(line[:320] + (" [LINE TRUNCATED]" if len(line) > 320 else ""))
    readable = "\n".join(body_lines)
    excerpt = head[:1200]
    if readable:
        excerpt += "\n\n[BODY EXCERPT]\n" + readable[:600]
    if len(text) > len(excerpt):
        excerpt += f"\n\n[TRUNCATED FOR READABILITY - complete {len(text)} character original retained in the evidence bundle]"
    return excerpt[:limit]


def _pdf_cell(value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(_pdf_text(value, 2000), style)


def _pdf_page_number(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#607481"))
    canvas.drawRightString(A4[0] - 42, 24, f"Page {document.page}")
    canvas.restoreState()


def _doc_page_number(document: Document) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _curated_findings(findings: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    """Separate risks, observations and semantically duplicated observations."""
    risks = [item for item in findings if item.get("severity") != "info"]
    observations = [item for item in findings if item.get("severity") == "info"]
    has_specific_header_risk = any(
        (item.get("evidence") or {}).get("evidence_type") == "http_response_header_absence"
        for item in risks
    )
    suppressed: list[dict] = []
    kept: list[dict] = []
    for item in observations:
        title = str(item.get("title") or "").strip().lower()
        if has_specific_header_risk and title == "http missing security headers":
            suppressed.append(item)
        else:
            kept.append(item)
    return risks, kept, suppressed


def _overall_risk(findings: list[dict]) -> str:
    for severity in SEVERITIES[:-1]:
        if any(item.get("severity") == severity and (item.get("validation_status") or "candidate") != "rejected" for item in findings):
            return severity.upper()
    return "NONE IDENTIFIED"


def _duration(stage: dict) -> str:
    started, finished = stage.get("started_at"), stage.get("finished_at")
    if started and finished:
        seconds = max(0, int((finished - started).total_seconds()))
        return f"{seconds // 60}m {seconds % 60}s"
    return "-"


def _load_artifact_bytes(item: dict | None) -> bytes | None:
    if not item:
        return None
    try:
        content = artifact_path(item["storage_key"]).read_bytes()
        return content if hashlib.sha256(content).hexdigest() == item["sha256"] else None
    except (OSError, KeyError):
        return None


def _source_payload(finding: dict, artifacts: dict[str, dict]) -> dict | None:
    artifact_id = _source_artifact_id(finding)
    content = _load_artifact_bytes(artifacts.get(artifact_id))
    if not content:
        return None
    try:
        payload = json.loads(content)
        return payload if isinstance(payload, dict) else None
    except (UnicodeDecodeError, json.JSONDecodeError):
        return None


def _source_artifact_id(finding: dict) -> str:
    return str((finding.get("evidence") or {}).get("source_artifact_id") or finding.get("source_artifact_id") or "")


def _shared_exchange_key(finding: dict, source_payload: dict | None) -> str | None:
    evidence = finding.get("evidence") or {}
    if evidence.get("request") or evidence.get("response") or not source_payload:
        return None
    return _source_artifact_id(finding) or None


def _captured_exchange(payload: dict | None) -> list[tuple[str, str]]:
    if not payload:
        return []
    payload = redact_object(payload)
    request = payload.get("request") if isinstance(payload.get("request"), dict) else {}
    rows: list[tuple[str, str]] = []
    for label, value in (
        ("Request Method", request.get("method")),
        ("Requested URL", request.get("url") or payload.get("requested_url")),
        ("Request Headers", request.get("headers")),
        ("Final URL", payload.get("final_url")),
        ("Response Status", payload.get("status")),
        ("Response Headers", payload.get("headers")),
        ("Response Body SHA-256", payload.get("body_sha256")),
        ("Response Body Truncated", payload.get("body_truncated")),
        ("Captured At", _format_timestamp(payload.get("observed_at")) if payload.get("observed_at") else None),
    ):
        if value not in (None, "", [], {}):
            rows.append((label, _evidence_value(value, 2400)))
    return rows


def _http_evidence_sections(evidence: dict, source_payload: dict | None = None) -> dict:
    evidence = redact_object(evidence)
    source_payload = redact_object(source_payload) if source_payload else None
    request = _protocol_text(evidence.get("request"))
    response = _protocol_text(evidence.get("response"))
    metadata = []
    preferred = (
        "template_id", "matcher_name", "matched_at", "requested_url", "status",
        "observed_at",
    )
    excluded = {
        "request", "response", "references", "requires_screenshot",
        "screenshot_artifact_id", "screenshot_sha256", "screenshot_role",
        "source_artifact_id", "source_sha256", "coverage_scope", "coverage_target",
    }
    ordered_keys = [key for key in preferred if key in evidence]
    ordered_keys.extend(sorted(key for key in evidence if key not in excluded and key not in ordered_keys))
    for key in ordered_keys:
        value = evidence.get(key)
        if value not in (None, "", [], {}):
            recorded = _format_timestamp(value) if key.endswith("_at") else _evidence_value(value)
            metadata.append((_evidence_label(key), recorded))
    return {
        "metadata": metadata,
        "request": request,
        "request_line": request.splitlines()[0] if request else None,
        "request_breakdown": _protocol_breakdown(evidence.get("request")),
        "response": response,
        "response_line": response.splitlines()[0] if response else None,
        "response_breakdown": _protocol_breakdown(evidence.get("response")),
        "captured_exchange": _captured_exchange(source_payload),
    }


def _protocol_breakdown(value: object) -> dict | None:
    text = redact_text(str(value or "")).replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    if not text:
        return None
    head, _, body = text.partition("\n\n")
    lines = [line for line in head.splitlines() if line.strip()]
    if not lines:
        return None
    headers: list[tuple[str, str]] = []
    for line in lines[1:]:
        if ":" in line:
            name, raw_value = line.split(":", 1)
            headers.append((_safe_text(name.strip(), 120), _safe_text(raw_value.strip(), 240)))
        else:
            headers.append(("Unparsed Header Line", _safe_text(line, 240)))
    body_preview = ""
    if body:
        body_preview = _safe_text(body.strip(), 500)
        if len(body.strip()) > 500:
            body_preview += " [TRUNCATED FOR READABILITY]"
    return {
        "start_line": _safe_text(lines[0], 240),
        "header_count": len(headers),
        "headers": headers,
        "body_present": bool(body.strip()),
        "body_preview": body_preview or None,
    }


def _finding_traceability_rows(finding: dict, artifacts: dict[str, dict]) -> list[list[str]]:
    evidence = finding.get("evidence") or {}
    source_artifact_id = _source_artifact_id(finding) or "-"
    source_artifact = artifacts.get(source_artifact_id)
    screenshot_artifact_id = _screenshot_artifact_id(finding) or "-"
    screenshot_artifact = artifacts.get(screenshot_artifact_id)
    screenshot_scope = _screenshot_scope_error(finding, artifacts)
    coverage_scope = evidence.get("coverage_scope") or "profile_observation"
    rows = [
        ["Finding record ID", str(finding.get("id") or "-")],
        ["Finding assurance", str(finding.get("validation_status") or "candidate").upper()],
        ["Evidence integrity", str(finding.get("evidence_integrity") or "not-reviewed").upper()],
        ["Evidence type", _safe_text(evidence.get("evidence_type") or evidence.get("template_id") or "scanner_observation", 120)],
        ["Source artifact ID", source_artifact_id],
        ["Source artifact SHA-256", str(evidence.get("source_sha256") or "-")],
        ["Source artifact kind", str((source_artifact or {}).get("kind") or "-")],
        ["Coverage scope", _safe_text(coverage_scope, 120)],
    ]
    recorded_target = str(finding.get("target") or "")
    if recorded_target and _display_url(recorded_target) != recorded_target:
        rows.append(["Raw scanner target", recorded_target])
    if screenshot_artifact_id != "-":
        rows.extend([
            ["Screenshot artifact ID", screenshot_artifact_id],
            ["Screenshot SHA-256", str(evidence.get("screenshot_sha256") or "-")],
            ["Screenshot final URL", str((screenshot_artifact or {}).get("metadata", {}).get("final_url") or "-")],
            ["Screenshot admissibility", "EXCLUDED - out of scope" if screenshot_scope else "ADMISSIBLE"],
        ])
    return rows


def _doc_technical_evidence(
    document: Document,
    finding: dict,
    artifacts: dict[str, dict],
    *,
    protocol_exhibit: str | None = None,
    shared_protocol_exhibit: str | None = None,
) -> None:
    view = _http_evidence_sections(finding.get("evidence") or {}, _source_payload(finding, artifacts))
    if view["metadata"]:
        _doc_table(document, ["Evidence Attribute", "Recorded Value"], view["metadata"])
    if view["request_breakdown"]:
        request_view = view["request_breakdown"]
        _doc_heading(document, "HTTP Request Summary", 2)
        _doc_table(document, ["Recorded Element", "Value"], [
            ("Start line", request_view["start_line"]),
            ("Headers recorded", request_view["header_count"]),
            ("Body present", "Yes" if request_view["body_present"] else "No"),
        ])
        if request_view["headers"]:
            _doc_table(document, ["Request Header", "Recorded Value"], request_view["headers"])
        if request_view["body_preview"]:
            _doc_body(document, f"Request body preview: {request_view['body_preview']}")
    if view["request"]:
        _doc_heading(document, "HTTP Request Evidence", 2)
        _doc_body(document, f"Exact scanner request. Start line: {view['request_line']}")
        _doc_protocol(document, view["request"])
    if view["response_breakdown"]:
        response_view = view["response_breakdown"]
        _doc_heading(document, "HTTP Response Summary", 2)
        _doc_table(document, ["Recorded Element", "Value"], [
            ("Status line", response_view["start_line"]),
            ("Headers recorded", response_view["header_count"]),
            ("Body present", "Yes" if response_view["body_present"] else "No"),
        ])
        if response_view["headers"]:
            _doc_table(document, ["Response Header", "Recorded Value"], response_view["headers"])
        if response_view["body_preview"]:
            _doc_body(document, f"Response body preview: {response_view['body_preview']}")
    if view["response"]:
        _doc_heading(document, "HTTP Response Evidence", 2)
        _doc_body(document, f"Exact target response associated with the match. Status line: {view['response_line']}")
        _doc_protocol(document, view["response"])
    if view["captured_exchange"]:
        if shared_protocol_exhibit:
            _doc_heading(document, "Shared Protocol Evidence", 2)
            _doc_body(document, f"See Protocol Exhibit {shared_protocol_exhibit}. This finding is a distinct condition derived from the same hash-verified HTTP exchange; only the finding-specific observation above changes.")
        else:
            title = f"Protocol Evidence Exhibit {protocol_exhibit}" if protocol_exhibit else "Captured Request and Response Metadata"
            _doc_heading(document, title, 2)
            _doc_body(document, "The following values were read from the original hash-verified source artifact. They are captured structured evidence, not a reconstructed wire exchange. Credential-bearing header values are redacted in the client report.")
            _doc_table(document, ["Captured Field", "Recorded Value"], view["captured_exchange"])
    if not view["request"] and not view["response"] and not view["captured_exchange"] and not shared_protocol_exhibit:
        _doc_body(document, "No readable protocol exchange was retained. Reliance is limited to the structured metadata and source-artifact hash recorded above.")


def _finding_reproduction(finding: dict) -> list[str]:
    evidence = finding.get("evidence") or {}
    template = evidence.get("template_id")
    source_hash = evidence.get("source_sha256") or "recorded in the evidence manifest"
    if evidence.get("evidence_type") == "http_response_header_absence":
        header = str(evidence.get("header_name") or "the named header")
        url = _display_url(evidence.get("requested_url") or finding["target"])
        return [
            f"Precondition: confirm the target is in the approved scope and perform only a read-only request to {url}.",
            f"Capture a fresh response without following redirects: curl -sS --max-redirs 0 -D response-headers.txt -o response-body.html \"{url}\".",
            f"Review response-headers.txt case-insensitively for a header named {header}. Expected result: no {header}: header line is present on the recorded response.",
            f"Record the response status, capture time, command, and full headers with the retest evidence. Do not compare a new response byte-for-byte with historical output because time-varying headers may differ.",
            f"For chain-of-custody, verify the original source artifact SHA-256 {source_hash} from the evidence bundle; that hash identifies the observation supporting this finding.",
        ]
    if str(template or "").lower() == "phpinfo-files":
        target = _display_url(finding["target"])
        return [
            f"From an authorized session, issue GET {urlsplit(target).path or '/'} to {target}.",
            "Confirm the response is successful and contains PHPInfo-specific markers such as 'PHP Version' or 'phpinfo()'.",
            "Record the complete response and an original browser capture without exposing credentials or unrelated sensitive values.",
            f"Compare the result with source artifact SHA-256 {source_hash} and the screenshot hash recorded in Technical Evidence.",
        ]
    steps = [
        f"Confirm the target {_display_url(finding['target'])} is still within the approved scope and use an authorized read-only test session.",
        f"Record the assessment profile, current timestamp, and exact URL before replaying {template or 'the documented assessment check'}.",
        f"Repeat the same non-destructive request or navigation path used by the platform and preserve the raw response, headers, and page state that support the observation.",
        f"Validate the expected condition described in this finding: {finding['description']}",
        f"Compare the replay result with source artifact SHA-256 {evidence.get('source_sha256') or 'recorded in the evidence manifest'} and retain any new evidence with its own timestamp and hash.",
        "If visual state is relevant, preserve an original screenshot from the reproduced step and tie it to the exact affected route rather than a later navigation state.",
    ]
    return steps


def _finding_retest(finding: dict) -> str:
    evidence = finding.get("evidence") or {}
    if evidence.get("evidence_type") == "http_response_header_absence":
        return f"Repeat the same authorized GET and confirm {evidence.get('header_name')} is present with the intended value on the affected response and representative authenticated routes. Verify no incompatible browser behavior was introduced."
    if str(evidence.get("template_id") or "").lower() == "phpinfo-files":
        return "Repeat the same GET and confirm the diagnostic resource returns 404/403 or an approved non-diagnostic response, then verify the PHPInfo markers are absent."
    return "Repeat the same bounded check after remediation and confirm the prior evidence condition is absent without introducing a regression."


def _assessment_perspective(context: dict) -> str:
    return "Authorized authenticated web" if context.get("authentication_configured") else "Unauthenticated external web"


def _load_screenshot(finding: dict, artifacts: dict[str, dict]) -> bytes | None:
    artifact_id = str((finding.get("evidence") or {}).get("screenshot_artifact_id") or "")
    item = artifacts.get(artifact_id)
    if not item or _screenshot_scope_error(finding, artifacts):
        return None
    return _load_artifact_bytes(item)


def _screenshot_scope_error(finding: dict, artifacts: dict[str, dict]) -> str | None:
    """Reject a browser image when its recorded final URL is not the finding origin."""
    evidence = finding.get("evidence") or {}
    artifact_id = str(evidence.get("screenshot_artifact_id") or "")
    item = artifacts.get(artifact_id)
    if not item:
        return None
    final_url = str((item.get("metadata") or {}).get("final_url") or "")
    target = str(finding.get("target") or evidence.get("coverage_target") or "")
    if not final_url or not target:
        return "The referenced browser capture lacks the final URL metadata required to verify its scope."
    expected, captured = urlsplit(target), urlsplit(final_url)
    expected_port = expected.port or (443 if expected.scheme == "https" else 80)
    captured_port = captured.port or (443 if captured.scheme == "https" else 80)
    if (
        captured.scheme != expected.scheme
        or captured.hostname != expected.hostname
        or captured_port != expected_port
    ):
        return "The referenced browser capture was excluded from this finding because its recorded final URL is outside the authorized origin. The original remains retained in the evidence bundle for audit."
    return None


def _finding_coverage_statement(finding: dict) -> str | None:
    evidence = finding.get("evidence") or {}
    if (
        evidence.get("coverage_scope") != "single_response"
        and evidence.get("evidence_type") != "http_response_header_absence"
    ):
        return None
    target = _display_url(evidence.get("coverage_target") or evidence.get("final_url") or finding.get("target"))
    return f"Coverage boundary: this configuration observation is confirmed only for the recorded response at {target}. It does not establish the same condition across other routes, hosts, or authenticated pages."


def _screenshot_artifact_id(finding: dict) -> str:
    return str((finding.get("evidence") or {}).get("screenshot_artifact_id") or "")


def _screenshot_caption(finding: dict, exhibit: str) -> str:
    evidence = finding.get("evidence") or {}
    if evidence.get("screenshot_role") == "contextual_browser_capture":
        return f"Contextual Screenshot Exhibit {exhibit}. Primary proof remains the hashed source response artifact; this image records the page state and visual context tied to the same route."
    return f"Original Screenshot Evidence Exhibit {exhibit}. The image is retained as an original capture and must be interpreted together with the source artifact hash and finding traceability table."


def _screenshot_segments(content: bytes) -> list[bytes]:
    """Split very tall captures into lossless, readable report segments."""
    with PILImage.open(io.BytesIO(content)) as image:
        width, height = image.size
        if height <= int(width * 1.5):
            return [content]
        segment_height = max(1, int(width * 0.72))
        segments: list[bytes] = []
        for top in range(0, height, segment_height):
            output = io.BytesIO()
            image.crop((0, top, width, min(top + segment_height, height))).save(output, format="PNG")
            segments.append(output.getvalue())
        return segments


def _screenshot_overview_size(content: bytes, max_width: float = 6.2, max_height: float = 4.2) -> tuple[float, float]:
    with PILImage.open(io.BytesIO(content)) as image:
        width, height = image.size
    scale = min(max_width / width, max_height / height)
    return width * scale, height * scale


def _screenshot_excerpt(content: bytes) -> bytes | None:
    """Return one readable viewport excerpt while preserving the original separately."""
    segments = _screenshot_segments(content)
    return segments[0] if len(segments) > 1 else None


def _artifact_summary(artifacts: list[dict]) -> list[list[object]]:
    counts: dict[str, tuple[int, int]] = {}
    for item in artifacts:
        count, size = counts.get(item["kind"], (0, 0))
        counts[item["kind"]] = (count + 1, size + int(item["size_bytes"]))
    return [[kind, count, size] for kind, (count, size) in sorted(counts.items())] or [["No retained artifacts", "-", "-"]]


def _terminal_artifacts(artifacts: list[dict]) -> list[dict]:
    return [item for item in artifacts if "terminal" in item["storage_key"].lower() and item["media_type"] == "image/png"]


def _terminal_exhibits(artifacts: list[dict], artifacts_by_id: dict[str, dict]) -> list[dict]:
    exhibits = []
    for item in _terminal_artifacts(artifacts):
        content = _load_artifact_bytes(item)
        if not content:
            continue
        metadata = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
        source_id = str(metadata.get("source_artifact_id") or "")
        source = artifacts_by_id.get(source_id)
        exhibits.append({
            "artifact": item,
            "content": content,
            "adapter": metadata.get("adapter") or Path(item["storage_key"]).stem.split("-terminal")[0],
            "transcript_sha256": metadata.get("source_transcript_sha256") or (source or {}).get("sha256") or "-",
            "provenance": "Hash-linked post-execution xterm capture showing the recorded command and final output excerpt. The complete retained transcript is the canonical machine-readable evidence for searching and independent hash verification.",
        })
    return exhibits


def render_docx(context: dict) -> bytes:
    assessment, scan = context["assessment"], context["scan"]
    findings, stages, artifacts = context["findings"], context["stages"], context["artifacts_by_id"]
    risk_findings, observations, suppressed = _curated_findings(findings)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    _doc_page_number(document)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("EXPOSURESCOPEX\n").bold = True
    title.add_run("Authorized Security Assessment Report").bold = True
    for run in title.runs:
        run.font.name = "Aptos Display"; run.font.size = Pt(24); run.font.color.rgb = RGBColor.from_string("16324F")
    subtitle = document.add_paragraph(f"{assessment['name']}\n{assessment['target']}\nEXECUTION {scan['status'].upper()} | {assessment['mode'].upper()} PROFILE")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("CONFIDENTIAL | AUTHORIZED SECURITY TESTING").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    _doc_heading(document, "Contents")
    _doc_table(document, ["Section Order", "Section"], [[index, section] for index, section in enumerate(REPORT_SECTIONS, start=1)])
    _doc_heading(document, "Document Control")
    _doc_heading(document, "Document History", 2)
    _doc_table(document, ["Version", "Issued", "Report Status", "Execution Status", "Change"], [["1.5", _format_timestamp(context["generated_at"]), context["report_status"].upper(), str(scan["status"]).upper(), "Corrected coverage accounting, secret redaction, URL normalization, and shared evidence exhibits"]])
    _doc_heading(document, "Point of Contact", 2)
    _doc_table(document, ["Organization", "Team", "Email"], [[settings().report_organization, settings().report_contact, getattr(settings(), "report_contact_email", "Not provided")]])
    _doc_heading(document, "Executive Summary")
    counts = {severity: sum(item["severity"] == severity for item in risk_findings) for severity in SEVERITIES[:-1]}
    assurance_counts = {
        state: sum((item.get("validation_status") or "candidate") == state for item in risk_findings)
        for state in ("confirmed", "candidate", "rejected", "inconclusive")
    }
    _doc_body(document, f"The authorized {assessment['mode']} assessment finished with operational status {scan['status']}. Overall technical risk is {_overall_risk(risk_findings)} based on {len(risk_findings)} distinct risk findings: {assurance_counts['confirmed']} confirmed, {assurance_counts['candidate']} candidates, {assurance_counts['rejected']} rejected, and {assurance_counts['inconclusive']} inconclusive. {len(observations)} non-risk security observations are reported separately and {len(suppressed)} duplicate observation was suppressed. All {len(context['artifacts'])} retained artifacts remain available in the evidence bundle.")
    _doc_table(document, ["Overall Risk", "Critical", "High", "Medium", "Low", "Observations"], [[_overall_risk(risk_findings), counts["critical"], counts["high"], counts["medium"], counts["low"], len(observations)]])
    _doc_body(document, f"Priority: remediate confirmed findings from highest to lowest severity, then independently review candidate observations. A COMPLETE execution status means every planned {assessment['mode']} stage reached a successful terminal state; it does not claim exhaustive penetration-test coverage.")
    _doc_heading(document, "Project Scope")
    _doc_table(document, ["Assessment", "Target", "Perspective", "Profile", "Plan", "Safety"], [[assessment["name"], assessment["target"], _assessment_perspective(context), assessment["mode"], scan.get("plan_version") or "-", "Non-exploitative"]])
    _doc_body(document, _scope_summary(context.get("scope")))
    _doc_body(document, f"Testing was restricted to the authorized origin and the immutable {assessment['mode']} plan. It did not include exploitation, destructive requests, persistence, credential attacks, source-code review, or manual business-logic abuse. Results therefore describe observed conditions within this profile, not all vulnerabilities that may exist.")
    _doc_heading(document, "Profiling")
    _doc_body(document, _profiling_summary_intro(context))
    _doc_heading(document, "Execution Summary", 2)
    _doc_table(document, ["Order", "Stage", "Required", "Status", "Duration", "Exception"], [[stage["position"] + 1, stage["adapter"], stage["required"], stage["status"], _duration(stage), stage.get("error_detail") or "-"] for stage in stages])
    _doc_body(document, _execution_exception_summary(stages))
    _doc_heading(document, "Execution Exceptions", 2)
    _doc_table(document, ["Order", "Stage", "Status", "Recorded Cause"], _execution_exception_rows(stages))
    _doc_heading(document, "Coverage Summary", 2)
    _doc_body(document, _coverage_summary_intro(context))
    _doc_heading(document, "Coverage by Test Family", 3)
    _doc_table(document, ["Family", "Required Cases", "Completed", "Exceptions", "Pending"], _coverage_family_rows(context.get("coverage") or []))
    for section_data in _profile_evidence_sections(context):
        _doc_heading(document, section_data["title"], 3)
        _doc_body(document, section_data["summary"])
        _doc_table(document, ["Measure", "Result"], section_data["rows"])
    _doc_heading(document, "Finding Index")
    _doc_body(document, "Use this index to triage findings before reading their full evidence, reproduction, remediation, and retest details. The evidence reference is the first 16 characters of the original source artifact SHA-256.")
    _doc_table(document, ["ID", "Severity", "Assurance", "Finding", "Affected Asset", "Evidence Ref"], _finding_index_rows(risk_findings))
    _doc_heading(document, "Detailed Findings")
    _doc_body(document, "The following distinct security weaknesses are normalized from retained evidence. Informational detections are intentionally excluded from the risk count.")

    labels = {"critical": "Critical Risk Findings", "high": "High Risk Findings", "medium": "Medium Risk Findings", "low": "Low Risk Findings", "info": "Informational Findings"}
    screenshot_exhibits: dict[str, str] = {}
    protocol_exhibits: dict[str, str] = {}
    for severity in SEVERITIES[:-1]:
        _doc_heading(document, labels[severity])
        group = [item for item in risk_findings if item["severity"] == severity]
        if not group:
            _doc_body(document, f"No {severity} findings were recorded by completed assessment stages.")
            continue
        for index, finding in enumerate(group, start=1):
            _doc_heading(document, f"{severity.upper()}-{index:03d}: {finding['title']}", 2)
            _doc_table(document, ["Severity", "Confidence", "Validation", "Affected Asset", "Finding ID"], [[severity.upper(), f"{finding['confidence']}%", (finding.get("validation_status") or "candidate").upper(), _display_url(finding["target"]), finding["id"]]])
            _doc_heading(document, "Description", 2); _doc_body(document, finding["description"])
            coverage = _finding_coverage_statement(finding)
            if coverage:
                _doc_body(document, coverage)
            _doc_heading(document, "Risk Classification", 2)
            _doc_table(document, ["CVSS v3.1", "Vector", "CWE", "OWASP"], _finding_taxonomy(finding))
            _doc_body(document, "Risk rationale: severity reflects the directly observed weakness and its plausible standalone consequence. Chained exploitation and unverified business impact are not included.")
            _doc_heading(document, "Business Impact", 2); _doc_body(document, _finding_business_impact(finding))
            _doc_heading(document, "Evidence Traceability", 2)
            _doc_table(document, ["Traceability Element", "Recorded Value"], _finding_traceability_rows(finding, artifacts))
            source_payload = _source_payload(finding, artifacts)
            shared_key = _shared_exchange_key(finding, source_payload)
            protocol_exhibit = None
            shared_protocol_exhibit = protocol_exhibits.get(shared_key or "")
            if shared_key and not shared_protocol_exhibit:
                protocol_exhibit = f"P-{len(protocol_exhibits) + 1:03d}"
                protocol_exhibits[shared_key] = protocol_exhibit
            _doc_heading(document, "Technical Evidence", 2)
            _doc_technical_evidence(
                document,
                finding,
                artifacts,
                protocol_exhibit=protocol_exhibit,
                shared_protocol_exhibit=shared_protocol_exhibit,
            )
            screenshot = _load_screenshot(finding, artifacts)
            if screenshot:
                artifact_id = _screenshot_artifact_id(finding)
                exhibit = screenshot_exhibits.get(artifact_id)
                if exhibit:
                    _doc_heading(document, "Shared Screenshot Evidence", 2)
                    _doc_body(document, f"See contextual Exhibit {exhibit}. The same original capture is referenced because this finding derives from the same recorded route and response. The unique condition for this finding is established by the traceability table and technical evidence above. SHA-256: {(finding.get('evidence') or {}).get('screenshot_sha256')}")
                else:
                    exhibit = f"S-{len(screenshot_exhibits) + 1:03d}"
                    screenshot_exhibits[artifact_id] = exhibit
                    _doc_heading(document, "Original Browser Evidence", 2)
                    _doc_body(document, _screenshot_caption(finding, exhibit))
                    width, height = _screenshot_overview_size(screenshot)
                    document.add_picture(io.BytesIO(screenshot), width=Inches(width), height=Inches(height))
                    excerpt = _screenshot_excerpt(screenshot)
                    if excerpt:
                        _doc_body(document, f"Exhibit {exhibit} readable viewport excerpt. The complete original remains above and in the evidence bundle.")
                        document.add_picture(io.BytesIO(excerpt), width=Inches(6.2))
                    _doc_body(document, f"Original capture SHA-256: {(finding.get('evidence') or {}).get('screenshot_sha256')}")
            elif scope_error := _screenshot_scope_error(finding, artifacts):
                _doc_body(document, scope_error)
            elif (finding.get("evidence") or {}).get("requires_screenshot"):
                _doc_body(document, "A screenshot was required by the originating check but no hash-verified image was available. Treat this finding as evidence-incomplete.")
            else:
                _doc_body(document, "A browser screenshot is not primary proof for this condition. Reliance is based on the hash-verified captured response evidence above.")
            _doc_heading(document, "Safe Reproduction Steps", 2)
            for step, value in enumerate(_finding_reproduction(finding), start=1):
                _doc_body(document, f"{step}. {value}")
            _doc_heading(document, "Remediation", 2); _doc_body(document, finding["remediation"])
            _doc_heading(document, "References", 2); _doc_body(document, "\n".join(_finding_references(finding)) or "No external reference was retained; consult the evidence manifest and originating scanner template.")
            _doc_heading(document, "Retest Criteria", 2); _doc_body(document, _finding_retest(finding))

    _doc_heading(document, "Informational Findings")
    _doc_body(document, "These detections provide asset and technology context but are not counted as vulnerabilities or included in the overall risk rating. Candidate observations require analyst review before client reliance.")
    if observations:
        _doc_table(document, ["Observation", "Assurance", "Asset", "Evidence Ref"], [[item["title"], (item.get("validation_status") or "candidate").upper(), _display_url(item["target"]), str((item.get("evidence") or {}).get("source_sha256") or "-")[:16]] for item in observations])
    else:
        _doc_body(document, "No distinct informational observations were retained.")
    if suppressed:
        _doc_body(document, f"Deduplication note: {len(suppressed)} generic observation was suppressed because its condition is represented by more specific risk findings.")

    _doc_heading(document, "Benchmark Evaluation")
    benchmark = context.get("benchmark")
    if benchmark:
        raw_score = benchmark["reported_findings_score"]
        score = benchmark["confirmed_findings_score"]
        eligibility = benchmark["release_eligibility"]
        quality_gate = score.get("quality_gate") or {}
        _doc_body(document, f"Benchmark scope is limited to {score['sample_size']} labelled cases applicable to the {score['profile']} profile in the pinned DVWA low-security fixture. These metrics do not measure complete DVWA or general web-application coverage.")
        _doc_table(document, ["Scoring view", "TP", "FP", "FN", "TN", "Precision", "Recall", "F1"], [
            ["Scanner detections before validation", raw_score["confusion"]["tp"], raw_score["confusion"]["fp"], raw_score["confusion"]["fn"], raw_score["confusion"]["tn"], f"{raw_score['precision']:.1%}", f"{raw_score['recall']:.1%}", f"{raw_score['f1']:.1%}"],
            ["Confirmed findings after validation", score["confusion"]["tp"], score["confusion"]["fp"], score["confusion"]["fn"], score["confusion"]["tn"], f"{score['precision']:.1%}", f"{score['recall']:.1%}", f"{score['f1']:.1%}"],
        ])
        _doc_table(document, ["Dataset", "Status", "Sample", "Release gate basis"], [[score["dataset_version"], "RELEASE" if eligibility["eligible"] else "DEVELOPMENT", score["sample_size"], "Confirmed findings after evidence validation"]])
        _doc_table(document, ["Measure", "Result"], [["Evidence completeness", f"{score['evidence_completeness']:.1%}"], ["Accuracy", f"{score['accuracy']:.1%}" if score["accuracy"] is not None else "Not computable - no negative controls"], ["Specificity", f"{score['specificity']:.1%}" if score["accuracy"] is not None else "Not computable - no negative controls"], ["True-positive cases", ", ".join(score["true_positive_case_ids"]) or "None"], ["False-positive observations", ", ".join(score["false_positive_observation_ids"]) or "None"], ["False-negative cases", ", ".join(score["false_negative_case_ids"]) or "None"], ["Evidence failures", ", ".join(score["evidence_failure_case_ids"]) or "None"]])
        _doc_body(document, "Scanner detections before validation measure raw sensitivity. Confirmed findings after validation are the only numbers used for release-gate quality claims.")
        if quality_gate:
            thresholds = quality_gate["thresholds"]
            _doc_table(document, ["Measured Quality Gate", "Result"], [["Eligibility", "PASS" if quality_gate["eligible"] else "FAIL"], ["Precision threshold", f">= {thresholds['precision']:.0%}"], ["Recall threshold", f">= {thresholds['recall']:.0%}"], ["Evidence completeness threshold", f">= {thresholds['evidence_completeness']:.0%}"]])
            if quality_gate.get("reason"):
                _doc_body(document, quality_gate["reason"])
        if not eligibility["eligible"]:
            _doc_body(document, eligibility["reason"])
    else:
        _doc_body(document, "No versioned benchmark was attributed to this target. Accuracy, precision, recall, and F1 are therefore not claimed by this report.")

    _doc_heading(document, "Appendix A: Engagement Methodology")
    _doc_table(document, ["Phase", "Purpose", "Evidence Produced"], [["1. Authorization and scope", "Validate exact origin, authorization acknowledgement, and safety boundary.", "Scope manifest and preflight transcript"], ["2. Attack-surface profiling", "Establish HTTP reachability, services, TLS posture, and authenticated routes.", "HTTP, service, TLS, crawl, and browser artifacts"], ["3. Deterministic assessment", "Run the immutable profile and pinned signed non-intrusive template inventory.", "Tool output, template inventory, and terminal transcripts"], ["4. Finding normalization", "Deduplicate observations and preserve source-to-finding provenance.", "Finding records and canonical observation keys"], ["5. Independent validation", "Verify hashes and apply safe replay or deterministic evidence oracles.", "Assurance decisions and validation replay"], ["6. Reporting", "Disclose coverage, exceptions, risk, evidence, and benchmark limitations.", "DOCX, PDF, and evidence ZIP"]])
    _doc_body(document, "Execution used leases, heartbeats, bounded concurrency, rate limits, and hard watchdog deadlines. No exploitation, persistence, destructive requests, credential attacks, or data extraction were performed. A successful stage means its declared bounded work completed; it does not imply that every possible vulnerability class was tested.")
    _doc_heading(document, "Appendix B: Risk Methodology")
    _doc_body(document, "Technical severity reflects the plausible consequence of the observed weakness. Confidence and assurance are separate: CONFIRMED requires intact evidence plus a deterministic oracle or safe independent replay; CANDIDATE indicates a scanner match that still requires analyst validation; REJECTED means independent evidence contradicted the claim; INCONCLUSIVE means evidence integrity or replay was insufficient. Informational technology detections do not affect overall risk. Business risk requires asset criticality, exposure and compensating-control context and is not fabricated by this report.")
    _doc_heading(document, "Appendix C: Testing Methodologies")
    mode = assessment["mode"].capitalize()
    _doc_body(document, _testing_methodology_appendix(context))
    _doc_heading(document, "Appendix D: Execution Coverage and Exceptions")
    _doc_body(document, _execution_exception_summary(stages))
    _doc_table(document, ["Order", "Stage", "Required", "Status", "Exception"], [[stage["position"] + 1, stage["adapter"], stage["required"], stage["status"], stage.get("error_detail") or "-"] for stage in stages])
    _doc_heading(document, "Appendix E: Evidence Integrity Manifest")
    _doc_body(document, "The evidence ZIP contains the complete machine-readable manifest and every retained original. The report uses excerpts only for readability; hashes always refer to the complete original bytes.")
    _doc_table(document, ["Artifact Kind", "Count", "Total Bytes"], _artifact_summary(context["artifacts"]))
    terminals = _terminal_artifacts(context["artifacts"])
    if terminals:
        _doc_heading(document, "Terminal Evidence Index", 2)
        _doc_table(document, ["Capture", "SHA-256 Prefix"], [[Path(item["storage_key"]).name, item["sha256"][:16]] for item in terminals])
        document.add_page_break()
        _doc_heading(document, "Terminal Evidence Exhibits", 2)
        for index, exhibit in enumerate(_terminal_exhibits(context["artifacts"], artifacts), start=1):
            if index > 1:
                document.add_page_break()
            item = exhibit["artifact"]
            _doc_heading(document, f"T-{index:03d}: {exhibit['adapter']}", 2)
            _doc_body(document, f"Screenshot SHA-256: {item['sha256']}\nTranscript SHA-256: {exhibit['transcript_sha256']}\nProvenance: {exhibit['provenance']}")
            width, height = _screenshot_overview_size(exhibit["content"], 6.2, 4.5)
            document.add_picture(io.BytesIO(exhibit["content"]), width=Inches(width), height=Inches(height))
    output = io.BytesIO(); document.save(output); return output.getvalue()


def render_pdf(context: dict) -> bytes:
    buffer = io.BytesIO()
    risk_findings, observations, suppressed = _curated_findings(context["findings"])
    counts = {severity: sum(item["severity"] == severity for item in risk_findings) for severity in SEVERITIES[:-1]}
    assurance_counts = {state: sum((item.get("validation_status") or "candidate") == state for item in risk_findings) for state in ("confirmed", "candidate", "rejected", "inconclusive")}
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=23, leading=27, textColor=colors.HexColor("#16324F"), alignment=TA_CENTER)
    heading = ParagraphStyle("ReportHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#16324F"), spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("ReportBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11)
    protocol = ParagraphStyle("ProtocolEvidence", parent=styles["Code"], fontName="Courier", fontSize=6.8, leading=8.5, leftIndent=10, spaceAfter=8)
    table_header = ParagraphStyle("EvidenceTableHeader", parent=body, fontName="Helvetica-Bold", textColor=colors.white)
    heading.keepWithNext = True
    styles["Heading2"].keepWithNext = True
    styles["Heading3"].keepWithNext = True
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=54, title=f"ExposureScopeX report - {context['assessment']['name']}")
    story = [Paragraph("EXPOSURESCOPEX", title), Spacer(1, 12), Paragraph("Authorized Security Assessment Report", title), Spacer(1, 20), Paragraph(_pdf_text(context["assessment"]["name"]), styles["Heading2"]), Paragraph(_pdf_text(context["assessment"]["target"]), body), Paragraph(f"EXECUTION {context['scan']['status'].upper()} | {context['assessment']['mode'].upper()} PROFILE", body), Spacer(1, 24), Paragraph("CONFIDENTIAL | AUTHORIZED SECURITY TESTING", body), PageBreak()]
    story += [Paragraph("Contents", heading), Table([[str(index), section] for index, section in enumerate(REPORT_SECTIONS, start=1)], colWidths=[0.65*inch, 6.0*inch], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),8),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#EAF2F7"))]))]
    simple_sections = [
        ("Executive Summary", f"The authorized {context['assessment']['mode']} assessment finished with operational status {context['scan']['status']}. Overall technical risk is {_overall_risk(risk_findings)} based on {len(risk_findings)} distinct risk findings: {assurance_counts['confirmed']} confirmed, {assurance_counts['candidate']} candidates, {assurance_counts['rejected']} rejected and {assurance_counts['inconclusive']} inconclusive. {len(observations)} non-risk observations are reported separately; {len(suppressed)} duplicate observation was suppressed. COMPLETE means all planned profile stages succeeded, not exhaustive penetration-test coverage."),
        ("Project Scope", f"Target: {context['assessment']['target']} | Profile: {context['assessment']['mode']} | Plan: {context['scan'].get('plan_version') or '-'} | {_assessment_perspective(context)} perspective | Non-exploitative. {_scope_summary(context.get('scope'))} Testing excluded exploitation, destructive requests, persistence, credential attacks, source review and manual business-logic abuse."),
        ("Profiling", _profiling_summary_intro(context)),
    ]
    history_rows = [
        [Paragraph(value, table_header) for value in ["Version", "Issued", "Report Status", "Execution Status", "Change"]],
        [_pdf_cell(value, body) for value in ["1.5", _format_timestamp(context["generated_at"]), context["report_status"].upper(), str(context["scan"]["status"]).upper(), "Corrected coverage accounting, secret redaction, URL normalization, and shared evidence exhibits"]],
    ]
    story += [Paragraph("Document Control", heading), Paragraph("Document History", styles["Heading3"]), Table(history_rows, colWidths=[.7*inch, 1.2*inch, .9*inch, 1.05*inch, 3.15*inch], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),4)])), Paragraph("Point of Contact", styles["Heading3"]), Table([["Organization", "Team", "Email"], [settings().report_organization, settings().report_contact, getattr(settings(), "report_contact_email", "Not provided")]], colWidths=[2.1*inch, 2.1*inch, 2.4*inch], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),4)]))]
    for name, value in simple_sections:
        story += [Paragraph(name, heading), Paragraph(_pdf_text(value), body)]
    story += [Paragraph("Risk Distribution", heading), Table([["Overall Risk", "Critical", "High", "Medium", "Low", "Observations"], [_overall_risk(risk_findings), counts["critical"], counts["high"], counts["medium"], counts["low"], len(observations)]], colWidths=[1.2*inch, .75*inch, .75*inch, .75*inch, .75*inch, 1.1*inch], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("ALIGN",(1,1),(-1,-1),"CENTER"),("FONTSIZE",(0,0),(-1,-1),8),("PADDING",(0,0),(-1,-1),5)])), Paragraph("Execution Summary", styles["Heading3"])]
    execution_rows = [[Paragraph(value, table_header) for value in ["No.", "Stage", "Req.", "Status", "Duration", "Exception"]]]
    execution_rows.extend([[item["position"] + 1, _pdf_text(item["adapter"]), "Yes" if item["required"] else "No", item["status"], _duration(item), _pdf_text(item.get("error_detail") or "-")] for item in context["stages"]])
    story.append(Table(execution_rows, colWidths=[.42*inch, 1.55*inch, .55*inch, .75*inch, .7*inch, 2.63*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),4)])))
    story.append(Paragraph(_pdf_text(_execution_exception_summary(context["stages"])), body))
    execution_exception_rows = [[Paragraph(value, table_header) for value in ["Order", "Stage", "Status", "Recorded Cause"]]]
    execution_exception_rows.extend([[_pdf_cell(row[0], body), _pdf_cell(row[1], body), _pdf_cell(row[2], body), _pdf_cell(row[3], body)] for row in _execution_exception_rows(context["stages"])])
    story += [Paragraph("Execution Exceptions", styles["Heading3"]), Table(execution_exception_rows, colWidths=[.55*inch, 1.85*inch, .95*inch, 3.25*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)]))]
    story += [Paragraph("Coverage Summary", styles["Heading3"]), Paragraph(_pdf_text(_coverage_summary_intro(context)), body), Paragraph("Coverage by Test Family", styles["Heading3"])]
    coverage_rows = [[Paragraph(value, table_header) for value in ["Family", "Required Cases", "Completed", "Exceptions", "Pending"]]]
    coverage_rows.extend([[_pdf_cell(value, body) for value in row] for row in _coverage_family_rows(context.get("coverage") or [])])
    story.append(Table(coverage_rows, colWidths=[2.05*inch, .95*inch, .85*inch, .85*inch, .85*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),4)])))
    for section_data in _profile_evidence_sections(context):
        section_rows = [[Paragraph(value, table_header) for value in ["Measure", "Result"]]]
        section_rows.extend([[_pdf_cell(value, body) for value in row] for row in section_data["rows"]])
        story += [
            Paragraph(section_data["title"], styles["Heading3"]),
            Paragraph(_pdf_text(section_data["summary"]), body),
            Table(section_rows, colWidths=[2.15*inch, 4.45*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),4)])),
        ]
    story += [Paragraph("Finding Index", heading), Paragraph("Use this index to triage findings before reading their full evidence, reproduction, remediation, and retest details.", body)]
    index_rows = [[Paragraph(value, table_header) for value in ["ID", "Severity", "Assurance", "Finding", "Asset", "Evidence Ref"]]]
    for row in _finding_index_rows(risk_findings):
        index_rows.append([row[0], *[_pdf_cell(value, body) for value in row[1:]]])
    story.append(Table(index_rows, colWidths=[.5*inch, .7*inch, .9*inch, 1.65*inch, 1.45*inch, 1.4*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),6.5),("PADDING",(0,0),(-1,-1),4)])))
    story += [Paragraph("Detailed Findings", heading), Paragraph("Distinct evidence-backed weaknesses are grouped by technical severity. Informational detections are excluded from the overall risk rating.", body)]
    labels = {"critical": "Critical Risk Findings", "high": "High Risk Findings", "medium": "Medium Risk Findings", "low": "Low Risk Findings"}
    screenshot_exhibits: dict[str, str] = {}
    protocol_exhibits: dict[str, str] = {}
    for severity in SEVERITIES[:-1]:
        story.append(Paragraph(labels[severity], heading))
        group = [item for item in risk_findings if item["severity"] == severity]
        if not group:
            story.append(Paragraph(f"No {severity} findings were recorded by completed stages.", body))
        for index, finding in enumerate(group, start=1):
            taxonomy = _finding_taxonomy(finding)[0]
            taxonomy_rows = [[Paragraph(value, table_header) for value in ["CVSS v3.1", "Vector", "CWE", "OWASP"]], [_pdf_cell(value, body) for value in taxonomy]]
            story += [Paragraph(f"{severity.upper()}-{index:03d}: {_pdf_text(finding['title'])}", styles["Heading2"]), Paragraph(f"Confidence: {finding['confidence']}% | Validation: {(finding.get('validation_status') or 'candidate').upper()} | Asset: {_pdf_text(_display_url(finding['target']))}", body), Paragraph(_pdf_text(finding["description"]), body)]
            coverage = _finding_coverage_statement(finding)
            if coverage:
                story.append(Paragraph(_pdf_text(coverage), body))
            trace_rows = [[Paragraph(value, table_header) for value in ["Traceability Element", "Recorded Value"]]]
            trace_rows.extend([[_pdf_cell(label, body), _pdf_cell(value, body)] for label, value in _finding_traceability_rows(finding, context["artifacts_by_id"])])
            story += [Paragraph("Risk Classification", styles["Heading3"]), Table(taxonomy_rows, colWidths=[.7*inch, 2.15*inch, 1.7*inch, 2.05*inch], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)])), Paragraph("Risk rationale: severity reflects the directly observed weakness and plausible standalone consequence; unverified exploit chains are excluded.", body), Paragraph("Business Impact", styles["Heading3"]), Paragraph(_pdf_text(_finding_business_impact(finding)), body), Paragraph("Evidence Traceability", styles["Heading3"]), Table(trace_rows, colWidths=[1.75*inch, 4.85*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)])), Paragraph("Technical Evidence", styles["Heading3"])]
            source_payload = _source_payload(finding, context["artifacts_by_id"])
            shared_key = _shared_exchange_key(finding, source_payload)
            protocol_exhibit = None
            shared_protocol_exhibit = protocol_exhibits.get(shared_key or "")
            if shared_key and not shared_protocol_exhibit:
                protocol_exhibit = f"P-{len(protocol_exhibits) + 1:03d}"
                protocol_exhibits[shared_key] = protocol_exhibit
            evidence_view = _http_evidence_sections(finding.get("evidence") or {}, source_payload)
            if evidence_view["metadata"]:
                evidence_rows = [[Paragraph("Evidence Attribute", table_header), Paragraph("Recorded Value", table_header)]]
                evidence_rows.extend([Paragraph(_pdf_text(label), body), Paragraph(_pdf_text(value), body)] for label, value in evidence_view["metadata"])
                story.append(Table(evidence_rows, colWidths=[1.45*inch, 5.15*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)])))
            if evidence_view["request_breakdown"]:
                request_view = evidence_view["request_breakdown"]
                request_rows = [[Paragraph("Recorded Element", table_header), Paragraph("Value", table_header)], [_pdf_cell("Start line", body), _pdf_cell(request_view["start_line"], body)], [_pdf_cell("Headers recorded", body), _pdf_cell(request_view["header_count"], body)], [_pdf_cell("Body present", body), _pdf_cell("Yes" if request_view["body_present"] else "No", body)]]
                story += [Paragraph("HTTP Request Summary", styles["Heading3"]), Table(request_rows, colWidths=[1.5*inch, 5.1*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)]))]
                if request_view["headers"]:
                    request_header_rows = [[Paragraph("Request Header", table_header), Paragraph("Recorded Value", table_header)]]
                    request_header_rows.extend([[_pdf_cell(label, body), _pdf_cell(value, body)] for label, value in request_view["headers"]])
                    story.append(Table(request_header_rows, colWidths=[2.1*inch, 4.5*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)])))
                if request_view["body_preview"]:
                    story.append(Paragraph(_pdf_text(f"Request body preview: {request_view['body_preview']}"), body))
            if evidence_view["request"]:
                story += [Paragraph("HTTP Request Evidence", styles["Heading3"]), Paragraph(_pdf_text(f"Exact scanner request. Start line: {evidence_view['request_line']}"), body), Preformatted(evidence_view["request"], protocol, maxLineLength=95)]
            if evidence_view["response_breakdown"]:
                response_view = evidence_view["response_breakdown"]
                response_rows = [[Paragraph("Recorded Element", table_header), Paragraph("Value", table_header)], [_pdf_cell("Status line", body), _pdf_cell(response_view["start_line"], body)], [_pdf_cell("Headers recorded", body), _pdf_cell(response_view["header_count"], body)], [_pdf_cell("Body present", body), _pdf_cell("Yes" if response_view["body_present"] else "No", body)]]
                story += [Paragraph("HTTP Response Summary", styles["Heading3"]), Table(response_rows, colWidths=[1.5*inch, 5.1*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)]))]
                if response_view["headers"]:
                    response_header_rows = [[Paragraph("Response Header", table_header), Paragraph("Recorded Value", table_header)]]
                    response_header_rows.extend([[_pdf_cell(label, body), _pdf_cell(value, body)] for label, value in response_view["headers"]])
                    story.append(Table(response_header_rows, colWidths=[2.1*inch, 4.5*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)])))
                if response_view["body_preview"]:
                    story.append(Paragraph(_pdf_text(f"Response body preview: {response_view['body_preview']}"), body))
            if evidence_view["response"]:
                story += [Paragraph("HTTP Response Evidence", styles["Heading3"]), Paragraph(_pdf_text(f"Exact target response associated with the match. Status line: {evidence_view['response_line']}"), body), Preformatted(evidence_view["response"], protocol, maxLineLength=95)]
            if evidence_view["captured_exchange"]:
                if shared_protocol_exhibit:
                    story += [Paragraph("Shared Protocol Evidence", styles["Heading3"]), Paragraph(_pdf_text(f"See Protocol Exhibit {shared_protocol_exhibit}. This finding is a distinct condition derived from the same hash-verified HTTP exchange; only the finding-specific observation above changes."), body)]
                else:
                    captured_rows = [[Paragraph("Captured Field", table_header), Paragraph("Recorded Value", table_header)]]
                    captured_rows.extend([[Paragraph(_pdf_text(label), body), Paragraph(_pdf_text(value), body)] for label, value in evidence_view["captured_exchange"]])
                    exhibit_title = f"Protocol Evidence Exhibit {protocol_exhibit}" if protocol_exhibit else "Captured Request and Response Metadata"
                    story += [Paragraph(exhibit_title, styles["Heading3"]), Paragraph("Values below are read from the original hash-verified source artifact and are not a reconstructed wire exchange. Credential-bearing header values are redacted in the client report.", body), Table(captured_rows, colWidths=[1.45*inch, 5.15*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)]))]
            if not evidence_view["request"] and not evidence_view["response"] and not evidence_view["captured_exchange"] and not shared_protocol_exhibit:
                story.append(Paragraph("No readable protocol exchange was retained. Reliance is limited to structured metadata and the source-artifact hash.", body))
            screenshot = _load_screenshot(finding, context["artifacts_by_id"])
            if screenshot:
                artifact_id = _screenshot_artifact_id(finding)
                exhibit = screenshot_exhibits.get(artifact_id)
                if exhibit:
                    story += [Paragraph("Shared Screenshot Evidence", styles["Heading3"]), Paragraph(_pdf_text(f"See contextual Exhibit {exhibit}. The same original capture is referenced because this finding derives from the same recorded route and response. The unique condition for this finding is established by the traceability table and technical evidence above. SHA-256: {(finding.get('evidence') or {}).get('screenshot_sha256')}"), body)]
                else:
                    exhibit = f"S-{len(screenshot_exhibits) + 1:03d}"
                    screenshot_exhibits[artifact_id] = exhibit
                    story += [Paragraph("Original Browser Evidence", styles["Heading3"]), Paragraph(_pdf_text(_screenshot_caption(finding, exhibit)), body)]
                    width, height = _screenshot_overview_size(screenshot, 6.4, 4.2)
                    story.append(Image(io.BytesIO(screenshot), width=width*inch, height=height*inch))
                    excerpt = _screenshot_excerpt(screenshot)
                    if excerpt:
                        story.append(Paragraph(f"Exhibit {exhibit} readable viewport excerpt. The complete original remains above and in the evidence bundle.", body))
                        reader = ImageReader(io.BytesIO(excerpt)); image_width, image_height = reader.getSize(); scale = 6.4*inch/image_width
                        story.append(Image(io.BytesIO(excerpt), width=image_width*scale, height=image_height*scale))
                    story.append(Paragraph(f"Original capture SHA-256: {(finding.get('evidence') or {}).get('screenshot_sha256')}", body))
            elif scope_error := _screenshot_scope_error(finding, context["artifacts_by_id"]):
                story.append(Paragraph(_pdf_text(scope_error), body))
            elif (finding.get("evidence") or {}).get("requires_screenshot"):
                story.append(Paragraph("A required screenshot was unavailable; treat this finding as evidence-incomplete.", body))
            else:
                story.append(Paragraph("A browser screenshot is not primary proof for this condition. Reliance is based on the hash-verified captured response evidence above.", body))
            story.append(Paragraph("Safe Reproduction Steps", styles["Heading3"]))
            for step, value in enumerate(_finding_reproduction(finding), start=1): story.append(Paragraph(f"{step}. {_pdf_text(value)}", body))
            story += [Paragraph("Remediation", styles["Heading3"]), Paragraph(_pdf_text(finding["remediation"]), body), Paragraph("References", styles["Heading3"]), Paragraph(_pdf_text("\n".join(_finding_references(finding)) or "No external reference was retained; consult the evidence manifest and originating scanner template."), body), Paragraph("Retest Criteria", styles["Heading3"]), Paragraph(_pdf_text(_finding_retest(finding)), body)]
    story += [Paragraph("Informational Findings", heading), Paragraph("These detections provide asset and technology context but are not counted as vulnerabilities or included in the overall risk rating. Candidate observations require analyst review.", body)]
    if observations:
        observation_rows = [[Paragraph(value, table_header) for value in ["Observation", "Assurance", "Asset", "Evidence Ref"]]]
        observation_rows.extend([[_pdf_cell(item["title"], body), _pdf_cell((item.get("validation_status") or "candidate").upper(), body), _pdf_cell(_display_url(item["target"]), body), str((item.get("evidence") or {}).get("source_sha256") or "-")[:16]] for item in observations])
        story.append(Table(observation_rows, colWidths=[1.7*inch, .9*inch, 1.7*inch, 2.3*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),4)])))
    if suppressed:
        story.append(Paragraph(f"Deduplication note: {len(suppressed)} generic observation was suppressed because specific risk findings represent the same condition.", body))
    story.append(Paragraph("Benchmark Evaluation", heading))
    benchmark = context.get("benchmark")
    if benchmark:
        raw_score = benchmark["reported_findings_score"]; score = benchmark["confirmed_findings_score"]; eligibility = benchmark["release_eligibility"]; quality_gate = score.get("quality_gate") or {}
        story.append(Paragraph(f"Benchmark scope is limited to {score['sample_size']} labelled cases applicable to the {score['profile']} profile in the pinned DVWA low-security fixture. It does not measure complete DVWA or general web-application coverage.", body))
        story.append(Table([["Scoring view", "TP", "FP", "FN", "TN", "Precision", "Recall", "F1"], ["Scanner detections before validation", raw_score["confusion"]["tp"], raw_score["confusion"]["fp"], raw_score["confusion"]["fn"], raw_score["confusion"]["tn"], f"{raw_score['precision']:.1%}", f"{raw_score['recall']:.1%}", f"{raw_score['f1']:.1%}"], ["Confirmed findings after validation", score["confusion"]["tp"], score["confusion"]["fp"], score["confusion"]["fn"], score["confusion"]["tn"], f"{score['precision']:.1%}", f"{score['recall']:.1%}", f"{score['f1']:.1%}"]], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),6.5),("PADDING",(0,0),(-1,-1),4)])))
        story.append(Table([["Dataset", "Status", "Sample", "Release gate basis"], [score["dataset_version"], "RELEASE" if eligibility["eligible"] else "DEVELOPMENT", score["sample_size"], "Confirmed findings after evidence validation"]], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),6.5),("PADDING",(0,0),(-1,-1),4)])))
        detail_rows = [["Measure", "Result"], ["Evidence completeness", f"{score['evidence_completeness']:.1%}"], ["Accuracy", f"{score['accuracy']:.1%}" if score["accuracy"] is not None else "Not computable - no negative controls"], ["Specificity", f"{score['specificity']:.1%}" if score["accuracy"] is not None else "Not computable - no negative controls"], ["True-positive cases", ", ".join(score["true_positive_case_ids"]) or "None"], ["False-positive observations", ", ".join(score["false_positive_observation_ids"]) or "None"], ["False-negative cases", ", ".join(score["false_negative_case_ids"]) or "None"], ["Evidence failures", ", ".join(score["evidence_failure_case_ids"]) or "None"]]
        story.append(Table([[Paragraph(_pdf_text(value), table_header) for value in detail_rows[0]], *[[_pdf_cell(row[0], body), _pdf_cell(row[1], body)] for row in detail_rows[1:]]], colWidths=[1.55*inch, 5.05*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)])))
        story.append(Paragraph("Scanner detections before validation measure raw sensitivity. Confirmed findings after validation are the only numbers used for release-gate quality claims.", body))
        if quality_gate:
            thresholds = quality_gate["thresholds"]
            quality_rows = [["Measured Quality Gate", "Result"], ["Eligibility", "PASS" if quality_gate["eligible"] else "FAIL"], ["Precision threshold", f">= {thresholds['precision']:.0%}"], ["Recall threshold", f">= {thresholds['recall']:.0%}"], ["Evidence completeness threshold", f">= {thresholds['evidence_completeness']:.0%}"]]
            story.append(Table([[Paragraph(_pdf_text(value), table_header) for value in quality_rows[0]], *[[_pdf_cell(row[0], body), _pdf_cell(row[1], body)] for row in quality_rows[1:]]], colWidths=[2.2*inch, 4.4*inch], repeatRows=1, style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)])))
            if quality_gate.get("reason"):
                story.append(Paragraph(_pdf_text(quality_gate["reason"]), body))
        if not eligibility["eligible"]: story.append(Paragraph(_pdf_text(eligibility["reason"]), body))
    else:
        story.append(Paragraph("No versioned benchmark was attributed to this target. Accuracy, precision, recall and F1 are not claimed.", body))
    appendices = [
        ("Appendix A: Engagement Methodology", "The engagement used six controlled phases: authorization and exact-origin scope validation; HTTP, service, TLS, and authenticated route profiling; deterministic execution of an immutable plan and pinned template inventory; immediate artifact ingestion and finding normalization; independent hash and safe-replay validation; and report finalization with explicit exceptions. Leases, heartbeats, rate limits, bounded concurrency, and hard watchdog deadlines controlled execution. Every stage retained its status and transcript. No exploitation, persistence, destructive requests, credential attacks, or data extraction were performed."),
        ("Appendix B: Risk Methodology", "Technical severity reflects plausible consequence. CONFIRMED requires intact evidence plus a deterministic oracle or safe replay. CANDIDATE requires analyst validation. REJECTED means independent evidence contradicted the claim. INCONCLUSIVE means evidence or replay was insufficient. Informational detections do not affect overall risk. Business risk requires client context."),
        ("Appendix C: Testing Methodologies", _testing_methodology_appendix(context)),
        ("Appendix D: Execution Coverage and Exceptions", _execution_exception_summary(context["stages"]) + " | " + " | ".join(f"{item['adapter']}={item['status']} ({item.get('error_detail') or 'no exception'})" for item in context["stages"])),
        ("Appendix E: Evidence Integrity Manifest", "The evidence ZIP contains the complete manifest and every retained original. Report excerpts do not replace originals and every displayed hash refers to complete source bytes."),
    ]
    for name, value in appendices: story += [Paragraph(name, heading), Paragraph(_pdf_text(value, 12000), body)]
    summary_rows = [["Artifact Kind", "Count", "Total Bytes"], *_artifact_summary(context["artifacts"])]
    story.append(Table(summary_rows, colWidths=[3.6*inch, 1*inch, 1.6*inch], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),4)])))
    terminals = _terminal_artifacts(context["artifacts"])
    if terminals:
        story.append(Paragraph("Terminal Evidence Index", styles["Heading3"]))
        terminal_rows = [["Capture", "SHA-256 Prefix"], *[[_pdf_cell(Path(item["storage_key"]).name, body), item["sha256"][:16]] for item in terminals]]
        story.append(Table(terminal_rows, colWidths=[2.7*inch, 3.9*inch], style=TableStyle([("GRID",(0,0),(-1,-1),.25,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),6.5),("PADDING",(0,0),(-1,-1),4)])))
        story.append(PageBreak())
        story.append(Paragraph("Terminal Evidence Exhibits", styles["Heading3"]))
        for index, exhibit in enumerate(_terminal_exhibits(context["artifacts"], context["artifacts_by_id"]), start=1):
            if index > 1:
                story.append(PageBreak())
            item = exhibit["artifact"]
            exhibit_story = [Paragraph(f"T-{index:03d}: {_pdf_text(exhibit['adapter'])}", styles["Heading3"]), Paragraph(_pdf_text(f"Screenshot SHA-256: {item['sha256']}\nTranscript SHA-256: {exhibit['transcript_sha256']}\nProvenance: {exhibit['provenance']}"), body)]
            width, height = _screenshot_overview_size(exhibit["content"], 6.4, 4.5)
            exhibit_story.append(Image(io.BytesIO(exhibit["content"]), width=width*inch, height=height*inch))
            story.append(KeepTogether(exhibit_story))
    doc.build(story, onFirstPage=_pdf_page_number, onLaterPages=_pdf_page_number); return buffer.getvalue()


def _benchmark_context(assessment: dict, scan: dict, findings: list[dict]) -> dict | None:
    manifest_path = Path(settings().benchmark_root) / "dvwa" / "manifest.json"
    if not manifest_path.is_file():
        return None
    manifest, truth = load_benchmark_manifest(manifest_path, variant="low", profile=assessment["mode"])
    if urlsplit(assessment["target"]).hostname not in manifest["fixture"].get("target_hosts", []):
        return None

    def observations(*, confirmed_only: bool) -> list[BenchmarkObservation]:
        result = []
        for finding in findings:
            if finding["severity"] == "info":
                continue
            if confirmed_only and not (
                finding.get("validation_status") == "confirmed"
                and finding.get("evidence_integrity") == "passed"
                and finding.get("oracle_status") == "passed"
            ):
                continue
            evidence = finding.get("evidence") or {}
            key = finding.get("canonical_observation_key") or canonical_observation_key(evidence)
            result.append(BenchmarkObservation(
                observation_id=str(finding["id"]),
                case_id=map_observation_key(
                    manifest, key=key, variant="low", profile=assessment["mode"]
                ),
                evidence_kinds=observation_evidence_kinds(evidence),
            ))
        return result

    return {
        "variant": "low",
        "release_eligibility": benchmark_release_eligibility(manifest),
        "reported_findings_score": score_vulnerability_benchmark(
            dataset_version=manifest["dataset_version"], profile=assessment["mode"],
            truth_cases=truth, observations=observations(confirmed_only=False),
        ),
        "confirmed_findings_score": score_vulnerability_benchmark(
            dataset_version=manifest["dataset_version"], profile=assessment["mode"],
            truth_cases=truth, observations=observations(confirmed_only=True),
        ),
    }


async def generate_scan_reports(scan_id, assessment_id, report_status: str, coverage: list[dict]) -> dict:
    assessment = dict(await pool().fetchrow("SELECT * FROM assessments WHERE id = $1", assessment_id))
    scope_row = await pool().fetchrow("SELECT scope FROM assessment_scopes WHERE assessment_id = $1", assessment_id)
    authentication_configured = await pool().fetchval(
        "SELECT EXISTS(SELECT 1 FROM assessment_authentication WHERE assessment_id = $1)", assessment_id
    )
    scan = dict(await pool().fetchrow("SELECT * FROM scans WHERE id = $1", scan_id))
    stages = [dict(row) for row in await pool().fetch("SELECT * FROM stage_runs WHERE scan_id = $1 ORDER BY position", scan_id)]
    canonical_coverage = [dict(row) for row in await pool().fetch(
        "SELECT * FROM scan_coverage WHERE scan_id = $1 ORDER BY created_at, case_id",
        scan_id,
    )]
    if canonical_coverage:
        coverage = canonical_coverage
    elif not coverage:
        coverage = stages
    findings = [dict(row) for row in await pool().fetch(
        """
        SELECT f.*, fa.canonical_observation_key, fa.validation_status,
               fa.evidence_integrity, fa.oracle_status, fa.validation_artifact_id,
               fa.rationale AS assurance_rationale, fa.validated_at
        FROM findings f LEFT JOIN finding_assurance fa ON fa.finding_id = f.id
        WHERE f.scan_id = $1 ORDER BY f.created_at
        """,
        scan_id,
    )]
    artifacts = [dict(row) for row in await pool().fetch(
        "SELECT * FROM artifacts WHERE scan_id = $1 AND kind NOT IN ('report_docx', 'report_pdf', 'evidence_bundle') ORDER BY captured_at",
        scan_id,
    )]
    generated_at = datetime.now(timezone.utc).isoformat()
    benchmark = _benchmark_context(assessment, scan, findings)
    context = {"assessment": assessment, "scope": dict(scope_row["scope"]) if scope_row else None, "authentication_configured": authentication_configured, "scan": scan, "stages": stages, "findings": findings, "artifacts": artifacts, "artifacts_by_id": {str(item["id"]): item for item in artifacts}, "report_status": report_status, "generated_at": generated_at, "benchmark": benchmark, "coverage": coverage}
    docx, pdf = await asyncio.gather(asyncio.to_thread(render_docx, context), asyncio.to_thread(render_pdf, context))
    report_stage = {"scan_id": scan_id, "id": None, "position": 98}
    docx_record = await persist_bytes(report_stage, docx, kind="report_docx", name="assessment-report.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", metadata={"report_status": report_status, "format": "docx"})
    pdf_record = await persist_bytes(report_stage, pdf, kind="report_pdf", name="assessment-report.pdf", media_type="application/pdf", metadata={"report_status": report_status, "format": "pdf"})
    artifacts = [dict(row) for row in await pool().fetch(
        "SELECT * FROM artifacts WHERE scan_id = $1 AND kind NOT IN ('report_docx', 'report_pdf', 'evidence_bundle') ORDER BY captured_at",
        scan_id,
    )]
    evidence_manifest = {"schema_version": "1.1", "scan_id": str(scan_id), "assessment_id": str(assessment_id), "generated_at": generated_at, "status": report_status, "execution_status": str(scan["status"]), "artifacts": [{"id": str(item["id"]), "storage_key": item["storage_key"], "kind": item["kind"], "media_type": item["media_type"], "sha256": item["sha256"], "size_bytes": item["size_bytes"], "captured_at": item["captured_at"].isoformat()} for item in artifacts]}
    bundle = io.BytesIO()
    with zipfile.ZipFile(bundle, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("evidence-manifest.json", json.dumps(evidence_manifest, indent=2))
        for item in artifacts:
            path = artifact_path(item["storage_key"])
            if path.is_file() and hashlib.sha256(path.read_bytes()).hexdigest() == item["sha256"]:
                archive.write(path, f"artifacts/{item['storage_key']}")
    bundle_record = await persist_bytes(report_stage, bundle.getvalue(), kind="evidence_bundle", name="evidence-bundle.zip", media_type="application/zip", metadata={"report_status": report_status, "format": "evidence"})
    risk_findings, observations, suppressed = _curated_findings(findings)
    return json_safe({"schema_version": "2.5", "scan_id": str(scan_id), "generated_at": generated_at, "status": report_status, "execution_status": str(scan["status"]), "finding_count": len(risk_findings), "observation_count": len(observations), "suppressed_duplicate_count": len(suppressed), "confirmed_finding_count": sum((item.get("validation_status") or "candidate") == "confirmed" for item in risk_findings), "candidate_finding_count": sum((item.get("validation_status") or "candidate") == "candidate" for item in risk_findings), "artifact_count": len(artifacts), "coverage": coverage, "benchmark": benchmark, "outputs": {"docx": {key: str(value) for key, value in docx_record.items() if key in {"id", "sha256", "storage_key", "size_bytes"}}, "pdf": {key: str(value) for key, value in pdf_record.items() if key in {"id", "sha256", "storage_key", "size_bytes"}}, "evidence": {key: str(value) for key, value in bundle_record.items() if key in {"id", "sha256", "storage_key", "size_bytes"}}}})
